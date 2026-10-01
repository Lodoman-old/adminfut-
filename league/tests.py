import ast
import os
import re
from datetime import date, datetime, timedelta, time
from io import BytesIO

from django.test import TestCase, Client
from django.utils import timezone
from django.core.management import call_command

from .models import Campo, Categoria, Equipo, Partido, Temporada, Jugador, Jornada, JugadorEquipo, MovimientoEquipo, SuspensionJugador, JugadorHerencia, AbandonoTemporada, Arbitro
from .cedula_service import procesar_cedula
from .reglas_movimientos import errores_movimiento_jugador, aplicar_movimiento_a_jugador
from .views import _hora_por_defecto_partido


class FixtureDescansoTest(TestCase):
    def setUp(self):
        self.cat = Categoria.objects.create(nombre="Libre", dias_juego=["SAB"])
        self.campo = Campo.objects.create(nombre="Campo 1", activo=True)
        self.equipos = [
            Equipo.objects.create(nombre=f"Equipo {i}", categoria=self.cat, activo=True)
            for i in range(5)
        ]
        self.fecha_inicio = date.today()

    def test_round_robin_impar_genera_descanso(self):
        t = Temporada(categoria=self.cat, nombre="Temp", fecha_inicio=self.fecha_inicio)
        fixture = t._generar_fixture_round_robin(self.equipos)

        self.assertEqual(len(fixture), 5)  # 5 jornadas con 5 equipos
        descansan = []
        for ronda in fixture:
            self.assertEqual(len(ronda), 2)  # (5-1)/2 partidos por jornada
            jugaron = []
            for local, visit in ronda:
                jugaron.extend([local.id, visit.id])
            descansan.append([e.id for e in self.equipos if e.id not in jugaron])

        # Cada equipo descansa exactamente una vez en la vuelta
        for d in descansan:
            self.assertEqual(len(d), 1)
        rests = [d[0] for d in descansan]
        self.assertEqual(len(set(rests)), 5)

        # Todos contra todos: cada pareja aparece exactamente una vez
        parejas = []
        for ronda in fixture:
            for local, visit in ronda:
                parejas.append(tuple(sorted([local.id, visit.id])))
        self.assertEqual(len(parejas), 10)
        self.assertEqual(len(set(parejas)), 10)

    def test_generar_rol_con_descanso_y_equipos_descansan(self):
        t = Temporada.objects.create(
            categoria=self.cat, nombre="Temp", fecha_inicio=self.fecha_inicio,
            tipo_rol="TODOS", vueltas=1,
        )
        t.generar_rol()

        jornadas = list(t.jornadas.order_by("numero"))
        self.assertEqual(len(jornadas), 5)

        for j in jornadas:
            partidos = Partido.objects.filter(jornada=j)
            self.assertEqual(partidos.count(), 2)
            jugaron = set(partidos.values_list("equipo_local_id", flat=True))
            jugaron.update(partidos.values_list("equipo_visitante_id", flat=True))
            descansan = t.equipos_descansan(j)
            self.assertEqual([e.id for e in descansan], [e.id for e in self.equipos if e.id not in jugaron])
            self.assertEqual(len(descansan), 1)

    def test_generar_rol_desde_jornada_inicial_genera_todo(self):
        from unittest import mock

        t = Temporada.objects.create(
            categoria=self.cat, nombre="Temp",
            fecha_inicio=date(2026, 5, 2),  # sábado, hace meses
            tipo_rol="TODOS", vueltas=1,
        )
        # Hoy = lunes 2026-05-18: el próximo día de juego (SAB) es 2026-05-23
        with mock.patch("django.utils.timezone.now", return_value=datetime(2026, 5, 18, 12, 0)):
            t.generar_rol(jornada_inicial=4)

        jornadas = list(t.jornadas.order_by("numero"))
        self.assertEqual(len(jornadas), 5)  # TODAS las jornadas se generan
        self.assertEqual([j.numero for j in jornadas], [1, 2, 3, 4, 5])

        # 1-3 con fechas pasadas desde fecha_inicio; 4-5 desde el próximo día de juego real
        fechas = [Partido.objects.filter(jornada=j).first().fecha_hora.date() for j in jornadas]
        self.assertEqual(fechas, [
            date(2026, 5, 2),
            date(2026, 5, 9),
            date(2026, 5, 16),
            date(2026, 5, 23),
            date(2026, 5, 30),
        ])


class HoraDefectoPartidoTest(TestCase):
    def test_hora_por_defecto_cicla_horarios_de_categoria(self):
        cat = Categoria.objects.create(nombre="Sabadito", horarios=["15:00", "17:00"])
        t = Temporada(categoria=cat)
        self.assertEqual(_hora_por_defecto_partido(t, 0), time(15, 0))
        self.assertEqual(_hora_por_defecto_partido(t, 1), time(17, 0))
        self.assertEqual(_hora_por_defecto_partido(t, 2), time(15, 0))

    def test_hora_por_defecto_sin_horarios_usa_12(self):
        cat = Categoria.objects.create(nombre="SinHoras")
        t = Temporada(categoria=cat)
        self.assertEqual(_hora_por_defecto_partido(t, 0), time(12, 0))


class PartidosPorJornadaAbandonoTest(TestCase):
    def test_equivale_al_rol_con_abandono(self):
        cat = Categoria.objects.create(nombre="ConBaja")
        temp = Temporada.objects.create(
            nombre="T", categoria=cat, fecha_inicio=date(2026, 1, 1)
        )
        equipos = []
        for i in range(6):
            equipos.append(Equipo.objects.create(nombre=f"E{i}", categoria=cat))
        # 5 equipos se quedan; el equipo 6 "se da de baja" (abandono)
        AbandonoTemporada.objects.create(temporada=temp, equipo=equipos[5])
        # El rol real incluye al de baja: 6 // 2 = 3 partidos por jornada
        self.assertEqual(temp.partidos_por_jornada(), 3)

    def test_par_de_equipos_sin_abandono(self):
        cat = Categoria.objects.create(nombre="SinBaja")
        temp = Temporada.objects.create(
            nombre="T2", categoria=cat, fecha_inicio=date(2026, 1, 1)
        )
        for i in range(6):
            Equipo.objects.create(nombre=f"E{i}", categoria=cat)
        self.assertEqual(temp.partidos_por_jornada(), 3)


class TablaPenalizacionDefaultTest(TestCase):
    """Cada derrota por default (equipo en abandono/mora) debe restar
    `puntos_default` de la tabla, una vez POR PARTIDO (regresión del bug
    que solo restaba la penalización en el primer partido)."""

    def setUp(self):
        self.cat = Categoria.objects.create(nombre="Default")
        self.temp = Temporada.objects.create(
            nombre="Temp", categoria=self.cat, fecha_inicio=date(2026, 1, 1),
            puntos_default=3,
        )
        self.e0, self.e1, self.e2, self.e3 = [
            Equipo.objects.create(nombre=f"E{i}", categoria=self.cat)
            for i in range(4)
        ]
        AbandonoTemporada.objects.create(temporada=self.temp, equipo=self.e0)

    def _partido(self, local, visitante):
        return Partido.objects.create(
            temporada=self.temp, equipo_local=local, equipo_visitante=visitante,
            fecha_hora=datetime(2026, 1, 1, 12, 0), estado="FIN",
        )

    def test_penaliza_una_vez_por_partido(self):
        self._partido(self.e0, self.e1)   # 0-1 (walkover)
        self._partido(self.e2, self.e0)   # 1-0
        self._partido(self.e0, self.e3)   # 0-1
        tabla = {r["nombre"]: r for r in self.temp.calcular_tabla()}
        self.assertEqual(tabla["E0 (Default)"]["pts"], -9)

    def test_ganador_walkover_pendiente(self):
        pendiente = Partido.objects.create(
            temporada=self.temp, equipo_local=self.e0, equipo_visitante=self.e1,
            fecha_hora=datetime(2026, 1, 10, 12, 0), estado="PEND",
        )
        self.assertEqual(pendiente.ganador_walkover, self.e1)

    def test_ganador_walkover_finalizado_default_local(self):
        p = Partido.objects.create(
            temporada=self.temp, equipo_local=self.e2, equipo_visitante=self.e3,
            fecha_hora=datetime(2026, 1, 10, 12, 0), estado="FIN",
            goles_local=1, goles_visitante=0, default_team="local",
        )
        self.assertEqual(p.ganador_walkover, self.e3)

    def test_walkover_libera_campo(self):
        """Un partido marcado como default (baja/alineación) no se juega:
        debe liberar el campo para dejar libre el bloque."""
        camp = Campo.objects.create(nombre="C", activo=True)
        p = Partido.objects.create(
            temporada=self.temp, equipo_local=self.e2, equipo_visitante=self.e3,
            campo=camp, fecha_hora=datetime(2026, 1, 10, 12, 0), estado="PEND",
        )
        p.default_team = "local"
        p.goles_local = 0
        p.goles_visitante = 1
        p.motivo_default = "Baja"
        p.estado = "FIN"
        p.save()
        p.refresh_from_db()
        self.assertIsNone(p.campo_id)
        self.assertIsNotNone(p.fecha_hora)

    def test_mora_libera_campo(self):
        """Walkover por mora/abandono también libera el campo."""
        camp = Campo.objects.create(nombre="C", activo=True)
        p = Partido.objects.create(
            temporada=self.temp, equipo_local=self.e0, equipo_visitante=self.e1,
            campo=camp, fecha_hora=datetime(2026, 1, 10, 12, 0), estado="PEND",
        )
        p.estado = "FIN"
        p.save()
        p.refresh_from_db()
        self.assertEqual((p.goles_local, p.goles_visitante), (0, 1))
        self.assertIsNone(p.campo_id)
        self.assertIsNotNone(p.fecha_hora)


class RolDespuesDeVueltaCompletaTest(TestCase):
    """Si las jornadas pasadas cubren una (o más) vueltas completas, las vueltas
    restantes (que repiten parejas) sí deben generarse (regresión del bug que
    dejaba 0 partidos cuando todas las parejas ya se habían enfrentado)."""

    def setUp(self):
        self.cat = Categoria.objects.create(nombre="Revueltas", dias_juego=["SAB"])
        self.campo = Campo.objects.create(nombre="Campo U", activo=True)
        self.equipos = [
            Equipo.objects.create(nombre=f"EQ{i}", categoria=self.cat, activo=True)
            for i in range(6)
        ]

    def _crear_vuelta_completa(self, temp, n_jornadas=5):
        """Registra a mano la primera vuelta completa (todas las parejas, una vez)."""
        eqs = temp._equipos_para_rol()[""]
        parejas = temp._pairings_robin_una_vuelta(eqs)
        self.assertEqual(len(parejas), 15)
        fecha = temp._proxima_fecha_juego()
        for i, (l_id, v_id) in enumerate(parejas):
            jornada, _ = Jornada.objects.get_or_create(
                temporada=temp, numero=(i // 3) + 1,
                defaults={"nombre": f"Jornada {(i // 3) + 1}"},
            )
            Partido.objects.create(
                temporada=temp, jornada=jornada,
                equipo_local_id=l_id, equipo_visitante_id=v_id,
                campo=self.campo, estado="PEND",
                fecha_hora=timezone.make_aware(datetime.combine(fecha, time(15, 0))),
            )

    def _crear_vuelta_completa_desordenada(self, temp):
        """Registra la primera vuelta pero con las parejas acomodadas en jornadas en
        un orden DISTINTO al que produce el algoritmo (reverso), para comprobar que
        las vueltas siguientes replican EL ORDEN REGISTRADO."""
        eqs = temp._equipos_para_rol()[""]
        parejas = temp._pairings_robin_una_vuelta(eqs)
        self.assertEqual(len(parejas), 15)
        fecha = temp._proxima_fecha_juego()
        for i, (l_id, v_id) in enumerate(reversed(parejas)):
            jornada, _ = Jornada.objects.get_or_create(
                temporada=temp, numero=((15 - 1 - i) // 3) + 1,
                defaults={"nombre": f"Jornada {(15 - 1 - i) // 3 + 1}"},
            )
            Partido.objects.create(
                temporada=temp, jornada=jornada,
                equipo_local_id=l_id, equipo_visitante_id=v_id,
                campo=self.campo, estado="PEND",
                fecha_hora=timezone.make_aware(datetime.combine(fecha, time(15, 0))),
            )

    def _parejas_de_jornada(self, temp, jn):
        out = []
        for l, v in temp.partidos.filter(jornada__numero=jn).values_list(
                "equipo_local_id", "equipo_visitante_id"):
            out.append((l, v))
        return sorted(out)

    def test_vueltas_siguientes_replican_el_orden_real_registrado(self):
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="Temp3", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=3,
        )
        self._crear_vuelta_completa_desordenada(temp)
        temp.generar_rol_respaldando_pasadas(jornada_inicial=6)

        for j in range(1, 6):
            base = temp.partidos.filter(jornada__numero=j)
            esperado = sorted(
                [(v, l) for l, v in base.values_list("equipo_local_id", "equipo_visitante_id")]
            )
            # Vuelta 2 (j+5): mismos enfrentamientos, localía invertida
            self.assertEqual(self._parejas_de_jornada(temp, j + 5), esperado, f"vuelta2 j{j}")
            # Vuelta 3 (j+10): mismos enfrentamientos que la vuelta 1
            self.assertEqual(
                self._parejas_de_jornada(temp, j + 10),
                sorted(base.values_list("equipo_local_id", "equipo_visitante_id")),
                f"vuelta3 j{j}",
            )

    def test_vuelta_entera_jugada_genera_las_vueltas_restantes(self):
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="Temp", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=3,
        )
        self._crear_vuelta_completa(temp)
        self.assertEqual(temp.partidos.count(), 15)

        temp.generar_rol_respaldando_pasadas(jornada_inicial=6)

        self.assertEqual(temp.partidos.count(), 45)  # 3 vueltas x15
        nums = sorted(temp.jornadas.values_list("numero", flat=True))
        self.assertEqual(nums, list(range(1, 16)))

        # Cada pareja quedó con exactamente 3 partidos (una por vuelta)
        from collections import Counter
        veces = Counter()
        for l, v in temp.partidos.values_list("equipo_local_id", "equipo_visitante_id"):
            veces[frozenset((l, v))] += 1
        self.assertEqual((set(veces.values())), {3})
        self.assertEqual(sum(veces.values()), 45)

        # Ningún equipo juega 2 veces en la misma jornada
        for j in Jornada.objects.filter(temporada=temp):
            pjs = temp.partidos.filter(jornada=j)
            ids = list(pjs.values_list("equipo_local_id", flat=True))
            ids += list(pjs.values_list("equipo_visitante_id", flat=True))
            self.assertEqual(len(ids), len(set(ids)))

    def test_sin_partidos_previos_sigue_generando_todas_las_vueltas(self):
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="Temp2", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=3,
        )
        temp.generar_rol_respaldando_pasadas(jornada_inicial=1)

        self.assertEqual(temp.partidos.count(), 45)
        self.assertEqual(sorted(temp.jornadas.values_list("numero", flat=True)), list(range(1, 16)))

    def test_fechas_futuras_empiezan_despues_de_la_ultima_pasada(self):
        """Regresión: el rol de jornadas futuras NO puede repetir la fecha de la
        última jornada pasada (antes partía de 'hoy', no de la última capturada)."""
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="FechasFuturas", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1,
        )
        eqs = temp._equipos_para_rol()[""]
        parejas = temp._pairings_robin_una_vuelta(eqs)
        fecha1 = date(2027, 1, 9)   # sábado
        fecha2 = date(2027, 1, 16)  # sábado, última jornada pasada
        for i, (l_id, v_id) in enumerate(parejas[:6]):
            jn = (i // 3) + 1
            fecha = fecha1 if jn == 1 else fecha2
            jornada, _ = Jornada.objects.get_or_create(
                temporada=temp, numero=jn,
                defaults={"nombre": f"Jornada {jn}"},
            )
            Partido.objects.create(
                temporada=temp, jornada=jornada,
                equipo_local_id=l_id, equipo_visitante_id=v_id,
                campo=self.campo, estado="PEND",
                fecha_hora=timezone.make_aware(datetime.combine(fecha, time(15, 0))),
            )

        temp.generar_rol_respaldando_pasadas(jornada_inicial=3)

        futuras = temp.partidos.filter(jornada__numero__gte=3)
        self.assertTrue(futuras.exists())
        for p in futuras:
            self.assertGreater(p.fecha_hora.date(), fecha2)

    def test_11_equipos_pasadas_parciales_rondas_completas_y_un_descanso(self):
        """11 equipos vueltas=1, capturadas 5 jornadas (orden arbitrario). El
        optimizador exhaustivo debe encontrar 6 rondas completas de 5, dando
        exactamente 11 jornadas con un descanso por equipo (una vuelta)."""
        cat = Categoria.objects.create(nombre="Primera", dias_juego=["SAB"])
        campo = Campo.objects.create(nombre="C", activo=True)
        eqs = [Equipo.objects.create(nombre=f"T{i:02d}", categoria=cat, activo=True)
               for i in range(11)]
        temp = Temporada.objects.create(
            categoria=cat, nombre="OnceR", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1,
        )
        ids = [e.id for e in eqs]
        parejas = temp._pairings_robin_una_vuelta(ids)
        self.assertEqual(len(parejas), 55)
        for i, (l, v) in enumerate(parejas[:25]):
            jn = i // 5 + 1
            j, _ = Jornada.objects.get_or_create(
                temporada=temp, numero=jn,
                defaults={"nombre": f"Jornada {jn}"},
            )
            Partido.objects.create(
                temporada=temp, jornada=j, equipo_local_id=l, equipo_visitante_id=v,
                campo=campo, estado="PEND",
                fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
            )

        temp.generar_rol_respaldando_pasadas(jornada_inicial=6)

        self.assertEqual(temp.partidos.count(), 55)
        nums = sorted(temp.jornadas.values_list("numero", flat=True))
        self.assertEqual(nums, list(range(1, 12)))
        for j in temp.jornadas.filter(numero__gte=6):
            self.assertEqual(
                temp.partidos.filter(jornada=j).count(), 5,
                f"jornada {j.numero} debe ser completa")

        from collections import Counter
        descansos = Counter()
        for j in temp.jornadas.order_by("numero"):
            pjs = temp.partidos.filter(jornada=j)
            ids_j = set(pjs.values_list("equipo_local_id", flat=True)) | set(pjs.values_list("equipo_visitante_id", flat=True))
            for e in ids:
                if e not in ids_j:
                    descansos[e] += 1
        self.assertEqual(set(descansos.values()), {1})

    def test_13_equipos_pasadas_arbitrarias_rondas_completas_y_un_descanso(self):
        """13 equipos vueltas=1, capturadas 5 jornadas en orden arbitrario (el
        caso real de Intermedia). El empaquetado por restarts debe dar 8 rondas
        completas de 6 -> 13 jornadas exactas, 1 descanso por equipo."""
        cat = Categoria.objects.create(nombre="Intermedia", dias_juego=["SAB"])
        campo = Campo.objects.create(nombre="C", activo=True)
        eqs = [Equipo.objects.create(nombre=f"T{i:02d}", categoria=cat, activo=True)
               for i in range(13)]
        temp = Temporada.objects.create(
            categoria=cat, nombre="TreceR", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1,
        )
        ids = [e.id for e in eqs]
        parejas = temp._pairings_robin_una_vuelta(ids)
        self.assertEqual(len(parejas), 78)
        import random as _random
        rondas = [parejas[i:i + 6] for i in range(0, 78, 6)]
        _random.Random(4040).shuffle(rondas)
        manual = [e for r in rondas[:5] for e in r]
        for i, (l, v) in enumerate(manual):
            jn = i // 6 + 1
            j, _ = Jornada.objects.get_or_create(
                temporada=temp, numero=jn,
                defaults={"nombre": f"Jornada {jn}"},
            )
            Partido.objects.create(
                temporada=temp, jornada=j, equipo_local_id=l, equipo_visitante_id=v,
                campo=campo, estado="PEND",
                fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
            )

        ok_p, errores_p = temp.validar_continuacion_rol()
        self.assertTrue(ok_p, errores_p)

        temp.generar_rol_respaldando_pasadas(jornada_inicial=6)

        self.assertEqual(temp.partidos.count(), 78)
        nums = sorted(temp.jornadas.values_list("numero", flat=True))
        self.assertEqual(nums, list(range(1, 14)))
        for j in temp.jornadas.filter(numero__gte=6):
            self.assertEqual(
                temp.partidos.filter(jornada=j).count(), 6,
                f"jornada {j.numero} debe ser completa")
        from collections import Counter
        veces = Counter()
        for l, v in temp.partidos.values_list("equipo_local_id", "equipo_visitante_id"):
            veces[frozenset((l, v))] += 1
        self.assertEqual(set(veces.values()), {1})
        descansos = Counter()
        for j in temp.jornadas.order_by("numero"):
            pjs = temp.partidos.filter(jornada=j)
            ids_j = set(pjs.values_list("equipo_local_id", flat=True)) | set(pjs.values_list("equipo_visitante_id", flat=True))
            for e in ids:
                if e not in ids_j:
                    descansos[e] += 1
        self.assertEqual(set(descansos.values()), {1})

    def test_validar_continuacion_detecta_descanso_repetido(self):
        """Si en las pasadas un equipo descansa dos veces, quedan más partidos
        pendientes que jornadas futuras: la validación debe devolver un error
        claro que impida regenerar y dejar un rol con rondas rotas."""
        cat = Categoria.objects.create(nombre="Imposible", dias_juego=["SAB"])
        campo = Campo.objects.create(nombre="C", activo=True)
        eqs = [Equipo.objects.create(nombre=f"T{i:02d}", categoria=cat, activo=True)
               for i in range(13)]
        temp = Temporada.objects.create(
            categoria=cat, nombre="TreceImposible", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1,
        )
        ids = [e.id for e in eqs]
        parejas = temp._pairings_robin_una_vuelta(ids)
        rondas = [parejas[i:i + 6] for i in range(0, 36, 6)]

        def swap_descanso(ronda, t00):
            presentes = {x for a, b in ronda for x in (a, b)}
            ausente = [t for t in ids if t not in presentes][0]
            out = []
            for a, b in ronda:
                if a == t00:
                    out.append((ausente, b))
                elif b == t00:
                    out.append((a, ausente))
                else:
                    out.append((a, b))
            return out

        t00 = ids[0]
        rondas[1] = swap_descanso(rondas[1], t00)   # T00 descansa aquí
        rondas[5] = swap_descanso(rondas[5], t00)   # y también aquí
        for jn, ronda in enumerate(rondas[:6], start=1):
            j, _ = Jornada.objects.get_or_create(
                temporada=temp, numero=jn,
                defaults={"nombre": f"Jornada {jn}"},
            )
            for l, v in ronda:
                Partido.objects.create(
                    temporada=temp, jornada=j, equipo_local_id=l, equipo_visitante_id=v,
                    campo=campo, estado="PEND",
                    fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
                )

        ok_p, errores_p = temp.validar_continuacion_rol()
        self.assertFalse(ok_p)
        self.assertTrue(any("T00" in e for e in errores_p), errores_p)

    def test_regenerar_desde_conserva_pasadas_y_rehace_desde_x(self):
        """El comando regenerar_desde borra solo jornadas >= X y conserva las
        capturadas a mano; el optimizador completa las futuras en rondas exactas
        aunque el greedy (orden de captura en el wizard) deje rondas cortas."""
        cat = Categoria.objects.create(nombre="Primera", dias_juego=["SAB"])
        campo = Campo.objects.create(nombre="C", activo=True)
        eqs = [Equipo.objects.create(nombre=f"T{i:02d}", categoria=cat, activo=True)
               for i in range(8)]
        temp = Temporada.objects.create(
            categoria=cat, nombre="OchoR", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1,
        )
        ids = [e.id for e in eqs]
        parejas = temp._pairings_robin_una_vuelta(ids)
        # Secuencia 2447: el residual es empaquetable en 4x4 pero el greedy
        # (por el orden de captura del wizard) dejaría 6 rondas -> 9 jornadas.
        import random as _random
        pool = list(parejas)
        _random.Random(2447).shuffle(pool)
        guardadas = [
            (l, v, i // 4 + 1)
            for i, (l, v) in enumerate(pool[:12])
        ]
        for l, v, jn in guardadas:
            j, _ = Jornada.objects.get_or_create(
                temporada=temp, numero=jn,
                defaults={"nombre": f"Jornada {jn}"},
            )
            Partido.objects.create(
                temporada=temp, jornada=j, equipo_local_id=l, equipo_visitante_id=v,
                campo=campo, estado="PEND",
                fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
            )

        # El greedy por sí solo dejaría 6 rondas (simulación del algoritmo viejo),
        # pero el optimizador exhaustivo ya empaqueta en 4x4 -> 7 jornadas.
        temp.generar_rol_respaldando_pasadas(jornada_inicial=4)
        self.assertEqual(temp.jornadas.count(), 7)

        call_command("regenerar_desde", temporada=temp.id, jornada_inicial=4,
                     force=True, verbosity=0)

        self.assertEqual(temp.jornadas.count(), 7)
        # Las 3 primeras conservan exactamente las parejas capturadas
        capturadas_actual = sorted(
            temp.partidos.filter(jornada__numero__lte=3)
            .values_list("equipo_local_id", "equipo_visitante_id"))
        self.assertEqual(
            capturadas_actual,
            sorted((l, v) for l, v, jn in guardadas),
        )
        # Futuras completas de 4
        for j in temp.jornadas.filter(numero__gte=4):
            self.assertEqual(
                temp.partidos.filter(jornada=j).count(), 4)
        self.assertEqual(
            set(Partido.objects.filter(temporada=temp)
                .values_list("jornada__numero", flat=True)),
            set(range(1, 8)),
        )

    def test_regenerar_rol_via_vista(self):
        """La vista regenerar_rol_temporada pide jornada_inicial por POST,
        conserva las pasadas y regenera desde X en adelante; sin permiso 403."""
        cat = Categoria.objects.create(nombre="PrimeraV", dias_juego=["SAB"])
        campo = Campo.objects.create(nombre="CV", activo=True)
        eqs = [Equipo.objects.create(nombre=f"U{i:02d}", categoria=cat, activo=True)
               for i in range(8)]
        temp = Temporada.objects.create(
            categoria=cat, nombre="OchoV", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1,
        )
        ids = [e.id for e in eqs]
        parejas = temp._pairings_robin_una_vuelta(ids)
        import random as _random
        pool = list(parejas)
        _random.Random(2447).shuffle(pool)
        guardadas = [(l, v, i // 4 + 1) for i, (l, v) in enumerate(pool[:12])]
        for l, v, jn in guardadas:
            j, _ = Jornada.objects.get_or_create(
                temporada=temp, numero=jn,
                defaults={"nombre": f"Jornada {jn}"},
            )
            Partido.objects.create(
                temporada=temp, jornada=j, equipo_local_id=l, equipo_visitante_id=v,
                campo=campo, estado="PEND",
                fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
            )

        from django.contrib.auth import get_user_model
        from django.urls import reverse

        # Primero se generarían las futuras (escenario real: rol ya creado)
        temp.generar_rol_respaldando_pasadas(jornada_inicial=4)
        self.assertEqual(temp.jornadas.count(), 7)

        normal = get_user_model().objects.create_user(username="sinpermv", password="p")
        self.client.force_login(normal)
        resp = self.client.post(
            reverse("regenerar_rol_temporada", args=[temp.id]),
            {"jornada_inicial": 4})
        self.assertEqual(resp.status_code, 403)
        self.client.logout()

        admin = get_user_model().objects.create_superuser(username="adminregen", password="p")
        self.client.force_login(admin)
        resp2 = self.client.post(
            reverse("regenerar_rol_temporada", args=[temp.id]),
            {"jornada_inicial": 4}, follow=True)
        self.assertEqual(resp2.status_code, 200)
        self.assertEqual(temp.jornadas.count(), 7)
        self.assertEqual(
            set(temp.partidos.filter(jornada__numero__lte=3)
                .values_list("equipo_local_id", "equipo_visitante_id")),
            set((l, v) for l, v, jn in guardadas),
        )
        for j in temp.jornadas.filter(numero__gte=4):
            self.assertEqual(temp.partidos.filter(jornada=j).count(), 4)


class ArbitroOpcionalTest(TestCase):
    """Los roles se generan con el árbitro en blanco (sin asignación
    automática) y la cédula no lo exige, pero lo guarda si se elige."""

    def setUp(self):
        self.cat = Categoria.objects.create(nombre="ArbOpCat", dias_juego=["SAB"])
        self.campo = Campo.objects.create(nombre="CampoArb", activo=True)
        self.arb = Arbitro.objects.create(nombre="Ref", apellido="Uno", activo=True)
        self.eqs = [Equipo.objects.create(nombre=f"T{i:02d}", categoria=self.cat, activo=True)
                    for i in range(6)]

    def test_generar_rol_deja_arbitro_en_blanco(self):
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="ArbBlanco", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1,
        )
        temp.generar_rol()
        partidos = Partido.objects.filter(temporada=temp)
        self.assertTrue(partidos.exists())
        self.assertFalse(partidos.exclude(arbitro__isnull=True).exists())

    def test_cedula_sin_arbitro_guarda_y_finaliza(self):
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="ArbCedula", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1, min_jugadores=0,
        )
        partido = Partido.objects.create(
            temporada=temp, equipo_local=self.eqs[0], equipo_visitante=self.eqs[1],
            campo=self.campo, estado="PEND",
            fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
        )
        r = procesar_cedula(partido, {"arbitro": "", "finalizar": "1"}, None)
        self.assertTrue(r["ok"], r["errors"])
        self.assertTrue(r["finalizado"])
        partido.refresh_from_db()
        self.assertEqual(partido.estado, "FIN")
        self.assertIsNone(partido.arbitro_id)

    def test_cedula_guarda_arbitro_cuando_se_elige(self):
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="ArbEleccion", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1, min_jugadores=0,
        )
        partido = Partido.objects.create(
            temporada=temp, equipo_local=self.eqs[0], equipo_visitante=self.eqs[1],
            campo=self.campo, estado="PEND",
            fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
        )
        r = procesar_cedula(partido, {"arbitro": str(self.arb.id), "finalizar": "1"}, None)
        self.assertTrue(r["ok"], r["errors"])
        partido.refresh_from_db()
        self.assertEqual(partido.arbitro_id, self.arb.id)

    def test_cedula_rechaza_arbitro_ya_ocupado_mismo_horario(self):
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="ArbConflicto", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1, min_jugadores=0,
        )
        hora = timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0)))
        ocupado = Partido.objects.create(
            temporada=temp, equipo_local=self.eqs[0], equipo_visitante=self.eqs[1],
            campo=self.campo, estado="PEND", fecha_hora=hora, arbitro=self.arb,
        )
        partido = Partido.objects.create(
            temporada=temp, equipo_local=self.eqs[2], equipo_visitante=self.eqs[3],
            campo=self.campo, estado="PEND", fecha_hora=hora,
        )
        r = procesar_cedula(partido, {"arbitro": str(self.arb.id), "finalizar": "1"}, None)
        self.assertFalse(r["ok"])
        self.assertTrue(any("ya está asignado" in e for e in r["errors"]), r["errors"])
        partido.refresh_from_db()
        self.assertIsNone(partido.arbitro_id)
        self.assertEqual(ocupado.arbitro_id, self.arb.id)

    def test_pdf_cedula_genera_sin_arbitro(self):
        from django.contrib.auth import get_user_model
        from django.urls import reverse
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="ArbPdf", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1, min_jugadores=0,
        )
        partido = Partido.objects.create(
            temporada=temp, equipo_local=self.eqs[0], equipo_visitante=self.eqs[1],
            campo=self.campo, estado="PEND",
            fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
        )
        for eq, prefijo in ((self.eqs[0], "LOC"), (self.eqs[1], "VIS")):
            for k in range(30):
                Jugador.objects.create(equipo=eq, nombre=f"{prefijo}{k} Cristóbal de los",
                                       apellido="Santos Hernández Gutiérrez de la Cruz",
                                       dorsal=k + 1, activo=True)
        u = get_user_model().objects.create_superuser(username="pdfadmin", password="p")
        self.client.force_login(u)
        resp = self.client.get(reverse("reporte_cedula_arbitral_pdf", args=[partido.id]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/pdf")
        self.assertGreater(len(resp.content), 500)

    def test_pdf_cedula_sin_campo_no_revienta(self):
        from django.contrib.auth import get_user_model
        from django.urls import reverse
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="ArbSinCampo", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1, min_jugadores=0,
        )
        partido = Partido.objects.create(
            temporada=temp, equipo_local=self.eqs[0], equipo_visitante=self.eqs[1],
            campo=None, estado="PEND",
            fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
        )
        for eq, prefijo in ((self.eqs[0], "LOC"), (self.eqs[1], "VIS")):
            for k in range(30):
                Jugador.objects.create(equipo=eq, nombre=f"{prefijo}{k} Cristóbal de los",
                                       apellido="Santos Hernández Gutiérrez de la Cruz",
                                       dorsal=k + 1, activo=True)
        u = get_user_model().objects.create_superuser(username="pdfsincampo", password="p")
        self.client.force_login(u)
        resp = self.client.get(reverse("reporte_cedula_arbitral_pdf", args=[partido.id]))
        self.assertEqual(resp.status_code, 200)
        import fitz
        doc = fitz.open(stream=resp.content, filetype="pdf")
        texto = "".join(p.get_text() for p in doc)
        self.assertIn("Por definir", texto)
        self.assertIn("Goles local", texto)
        self.assertIn("Firma del", texto)

    def test_pdf_cedula_firmas_y_goles_en_negro_aunque_haya_filas_vacias(self):
        from django.contrib.auth import get_user_model
        from django.urls import reverse
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="ArbBlanco", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1, min_jugadores=0,
        )
        partido = Partido.objects.create(
            temporada=temp, equipo_local=self.eqs[0], equipo_visitante=self.eqs[1],
            campo=self.campo, estado="PEND",
            fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
        )
        for eq, prefijo, n in ((self.eqs[0], "LOC", 30), (self.eqs[1], "VIS", 28)):
            for k in range(n):
                Jugador.objects.create(equipo=eq, nombre=f"{prefijo} LEOPOLDO ONITSED",
                                       apellido="SALDIVAR FLORES",
                                       dorsal=k + 1, activo=True)
        u = get_user_model().objects.create_superuser(username="pdfnegro", password="p")
        self.client.force_login(u)
        resp = self.client.get(reverse("reporte_cedula_arbitral_pdf", args=[partido.id]))
        self.assertEqual(resp.status_code, 200)
        import fitz
        doc = fitz.open(stream=resp.content, filetype="pdf")
        self.assertEqual(len(doc), 1, "La cédula debe caber en una sola hoja")
        blancos = []
        claves = ("Goles", "Capitan", "Firma", "Jugador", "Tit", "Camb", "A1", "A2", "Roja", "LOCAL", "VISITANTE")
        for page in doc:
            for b in page.get_text("dict")["blocks"]:
                if b.get("type") != 0:
                    continue
                for l in b.get("lines", []):
                    for s in l["spans"]:
                        if any(k in s["text"] for k in claves):
                            if s["color"] == 16777215:
                                blancos.append(s["text"])
        self.assertFalse(blancos, f"Texto invisible (color blanco): {blancos}")
        texto = "".join(p.get_text() for p in doc)
        self.assertIn("Goles local:", texto)
        self.assertIn("Firma del", texto)

    def test_pdf_cedula_muestra_marcador_cuando_finalizado(self):
        from django.contrib.auth import get_user_model
        from django.urls import reverse
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="ArbMarcador", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1, min_jugadores=0,
        )
        partido = Partido.objects.create(
            temporada=temp, equipo_local=self.eqs[0], equipo_visitante=self.eqs[1],
            campo=self.campo, estado="FIN", goles_local=2, goles_visitante=1,
            fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
        )
        for eq, prefijo in ((self.eqs[0], "LOC"), (self.eqs[1], "VIS")):
            for k in range(30):
                Jugador.objects.create(equipo=eq, nombre=f"{prefijo}{k} Cristóbal de los",
                                       apellido="Santos Hernández Gutiérrez de la Cruz",
                                       dorsal=k + 1, activo=True)
        u = get_user_model().objects.create_superuser(username="pdfmarcador", password="p")
        self.client.force_login(u)
        resp = self.client.get(reverse("reporte_cedula_arbitral_pdf", args=[partido.id]))
        self.assertEqual(resp.status_code, 200)
        import fitz
        doc = fitz.open(stream=resp.content, filetype="pdf")
        texto = "".join(p.get_text() for p in doc)
        self.assertIn("Goles local: 2", texto)
        self.assertIn("Goles visitante: 1", texto)

    def test_xlsx_cedula_genera_sin_arbitro(self):
        from django.contrib.auth import get_user_model
        from django.urls import reverse
        temp = Temporada.objects.create(
            categoria=self.cat, nombre="ArbXlsx", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1, min_jugadores=0,
        )
        partido = Partido.objects.create(
            temporada=temp, equipo_local=self.eqs[0], equipo_visitante=self.eqs[1],
            campo=self.campo, estado="PEND",
            fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
        )
        u = get_user_model().objects.create_superuser(username="xlsxadmin", password="p")
        self.client.force_login(u)
        resp = self.client.get(reverse("reporte_cedula_arbitral_xlsx", args=[partido.id]))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"].lower())
        import openpyxl
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        col_a = [ws.cell(row=r, column=1).value for r in range(7, ws.max_row + 1)]
        self.assertTrue(any(isinstance(v, str) and "LOCAL" in v for v in col_a), col_a)
        self.assertTrue(any(isinstance(v, str) and "VISITANTE" in v for v in col_a), col_a)


class LimiteCambiosTest(TestCase):
    def test_limite_cambios_efectivo(self):
        t = Temporada(cambios_permitidos=5)
        self.assertEqual(t.limite_cambios_efectivo(), 5)
        t.cambios_permitidos = 1
        self.assertEqual(t.limite_cambios_efectivo(), 1)
        t.cambios_permitidos = 0
        self.assertIsNone(t.limite_cambios_efectivo())
        t.cambios_permitidos = None
        self.assertIsNone(t.limite_cambios_efectivo())

    def test_guardar_cambios_cero_y_nulo(self):
        cat = Categoria.objects.create(nombre="LimiteCat")
        t0 = Temporada.objects.create(nombre="Sin cambios", categoria=cat, fecha_inicio=date(2026, 1, 1), cambios_permitidos=0)
        tn = Temporada.objects.create(nombre="Sin limite", categoria=cat, fecha_inicio=date(2026, 1, 1), cambios_permitidos=None)
        t5 = Temporada.objects.create(nombre="Con limite", categoria=cat, fecha_inicio=date(2026, 1, 1), cambios_permitidos=5)
        self.assertIsNone(t0.limite_cambios_efectivo())
        self.assertIsNone(tn.limite_cambios_efectivo())
        self.assertEqual(t5.limite_cambios_efectivo(), 5)


class CedulaInvitadoTest(TestCase):
    def setUp(self):
        self.cat = Categoria.objects.create(nombre="CedulaCat")
        self.campo = Campo.objects.create(nombre="Campo 1", activo=True)
        self.local = Equipo.objects.create(nombre="Local", categoria=self.cat, activo=True)
        self.visit = Equipo.objects.create(nombre="Visitante", categoria=self.cat, activo=True)
        self.fecha = timezone.make_aware(datetime(2026, 1, 1, 12, 0))

    def _partido(self, estado):
        return Partido.objects.create(
            equipo_local=self.local, equipo_visitante=self.visit,
            campo=self.campo, fecha_hora=self.fecha, estado=estado,
            goles_local=2, goles_visitante=1,
        )

    def test_invitado_ve_solo_cedulas_finalizadas(self):
        fin = self._partido("FIN")
        pend = self._partido("PEND")
        c = Client()
        r = c.get(f"/cedula-arbitral/{fin.id}/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Vista de solo lectura")
        c404 = Client(raise_request_exception=False)
        r2 = c404.get(f"/cedula-arbitral/{pend.id}/")
        self.assertEqual(r2.status_code, 404)
        from django.template.loader import render_to_string
        html = render_to_string("league/cedula_invitado_no_disponible.html", {"partido": pend})
        self.assertIn("Cédula no disponible", html)

    def test_home_invitado_usa_categoria_principal_de_cookie(self):
        otra = Categoria.objects.create(nombre="PrincipalCat", es_principal=True)
        c = Client()
        r = c.get("/")
        self.assertContains(r, f"Máximos Goleadores - {otra.nombre}")
        c.cookies["cat_preferida"] = str(self.cat.id)
        r2 = c.get("/")
        self.assertContains(r2, f"Máximos Goleadores - {self.cat.nombre}")


class PwaWebPushTest(TestCase):
    def test_endpoints_pwa_publicos(self):
        c = Client()
        r = c.get("/sw.js")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/javascript; charset=utf-8")
        self.assertContains(r, "showNotification")

        r = c.get("/manifest.webmanifest")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/manifest+json")
        self.assertIn("standalone", r.json().get("display", ""))

        r = c.get("/pwa-icon/icon-192.png")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "image/png")
        r = c.get("/pwa-icon/no-existe.png")
        self.assertEqual(r.status_code, 404)

        r = c.get("/api/vapid-public-key/")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json().get("public_key"))

    def test_registro_invitado_guarda_suscripcion_webpush(self):
        import json
        from .models import DeviceToken
        c = Client()
        r = c.post(
            "/api/register-guest/",
            data=json.dumps({
                "token": "web-111",
                "plataforma": "web",
                "device_id": "dev-pwa-001",
                "nombre": "Invitado PWA",
                "telefono": "6441112233",
                "email": "",
                "categorias": [],
                "subscripcion": {
                    "endpoint": "https://push.example.com/sub/abc123",
                    "p256dh": "BGhT9lP2T-k=",
                    "auth": "SxM9gQ==",
                },
            }),
            content_type="application/json",
        )
        self.assertEqual(r.status_code, 200)
        dt = DeviceToken.objects.get(device_id="dev-pwa-001")
        self.assertEqual(dt.plataforma, "pwa")
        self.assertEqual(dt.webpush_endpoint, "https://push.example.com/sub/abc123")
        self.assertEqual(dt.webpush_p256dh, "BGhT9lP2T-k=")
        self.assertEqual(dt.webpush_auth, "SxM9gQ==")
        self.assertTrue(dt.activo)

    def test_envio_webpush_desactiva_endpoint_410(self):
        from unittest import mock

        from pywebpush import WebPushException

        from .models import DeviceToken
        from .push import _enviar_webpush

        dt = DeviceToken.objects.create(
            token="w-410", device_id="dev-410",
            plataforma="pwa", activo=True,
            webpush_endpoint="https://push.example.com/sub/gone",
            webpush_p256dh="BGhT9lP2T-k=", webpush_auth="SxM9gQ==",
        )

        class Resp410:
            status_code = 410

        def fake_webpush(info, data, **kw):
            raise WebPushException("gone", Resp410())

        with mock.patch("pywebpush.webpush", side_effect=fake_webpush), \
                mock.patch("league.push._vapid_keys", return_value=("priv", "pub")):
            res = _enviar_webpush([dt], "T", "B", {"x": 1})

        self.assertEqual(res["success"], 0)
        self.assertEqual(res["failure"], 1)
        dt.refresh_from_db()
        self.assertFalse(dt.activo)


class CurpReglaMariaJoseTest(TestCase):
    def _form(self):
        from .forms import JugadorForm
        return JugadorForm()

    def test_inicial_nombre_curp_regla_maria_jose(self):
        f = self._form()
        self.assertEqual(f._inicial_nombre_curp("JOSE LUIS"), "L")
        self.assertEqual(f._inicial_nombre_curp("MARIA GUADALUPE"), "G")
        self.assertEqual(f._inicial_nombre_curp("MARIA DEL CARMEN"), "C")
        self.assertEqual(f._inicial_nombre_curp("JOSE MARIA"), "M")
        self.assertEqual(f._inicial_nombre_curp("JOSE"), "J")
        self.assertEqual(f._inicial_nombre_curp("LUIS"), "L")

    def test_inicial_nombre_curp_regla_abreviaturas(self):
        f = self._form()
        self.assertEqual(f._inicial_nombre_curp("J. CARMEN"), "C")
        self.assertEqual(f._inicial_nombre_curp("J. LUIS"), "L")
        self.assertEqual(f._inicial_nombre_curp("MA. GUADALUPE"), "G")
        self.assertEqual(f._inicial_nombre_curp("M. JOSE"), "J")
        self.assertEqual(f._inicial_nombre_curp("Jose Carmen"), "C")

    def test_validacion_acepta_regla_maria_jose(self):
        f = self._form()
        # José Luis García Hernández, 1990-01-01 → posición 4 = L (Luis)
        errs = f._validar_curp_contra_datos(
            "GAHL900101HDFRCR00", "José Luis", "García Hernández", date(1990, 1, 1)
        )
        self.assertEqual(errs, [])

    def _curp18(self, base17):
        from .forms import JugadorForm
        return base17 + str(JugadorForm._calcular_digito_curp(base17))

    def test_apellido_con_particula_paterno_deriva_del_nucleo(self):
        f = self._form()
        # 'De la Torre García' → de 'TORRE' (T,O) y 'GARCÍA' (G)
        curp = self._curp18("TOGP" + "900101" + "HDF" + "RCR" + "0")
        errs = f._validar_curp_contra_datos(
            curp, "Pedro", "De la Torre García", date(1990, 1, 1)
        )
        self.assertEqual(errs, [])

    def test_apellido_con_particula_materno_deriva_del_nucleo(self):
        f = self._form()
        # 'García De la Torre' → 'GARCÍA' (G,A) y materno 'TORRE' (T)
        curp = self._curp18("GATP" + "900101" + "HDF" + "RCR" + "0")
        errs = f._validar_curp_contra_datos(
            curp, "Pedro", "García De la Torre", date(1990, 1, 1)
        )
        self.assertEqual(errs, [])

    def test_apellido_unico_con_particula(self):
        f = self._form()
        # 'De León' (apellido único) → 'LEÓN' (L,E); sin materno usa 'X'
        curp = self._curp18("LEXJ" + "900101" + "HDF" + "RCR" + "0")
        errs = f._validar_curp_contra_datos(
            curp, "Juan", "De León", date(1990, 1, 1)
        )
        self.assertEqual(errs, [])

    def test_curp_real_de_la_cruz_acepta(self):
        f = self._form()
        # Caso real RENAPO: 'DE LA CRUZ FLORINDA' + 'FELIPE' → 'CUFF...'
        curp = self._curp18("CUFF" + "760301" + "HDF" + "RCR" + "0")
        errs = f._validar_curp_contra_datos(
            curp, "Felipe", "De la Cruz Florinda", date(1976, 3, 1)
        )
        self.assertEqual(errs, [])

    def test_apellido_con_particula_mal_capturado_si_marca_error(self):
        f = self._form()
        # Con la letra de la partícula ('D' en vez de 'T') debe reportar el error
        curp = self._curp18("DEGP" + "900101" + "HDF" + "RCR" + "0")
        errs = f._validar_curp_contra_datos(
            curp, "Pedro", "De la Torre García", date(1990, 1, 1)
        )
        self.assertTrue(any("apellido paterno" in e for e in errs))

    def test_validacion_rechaza_inicial_del_primer_nombre(self):
        f = self._form()
        # Con la inicial del primer nombre (J de José) debe marcar error explicando la regla
        errs = f._validar_curp_contra_datos(
            "GAHJ900101HDFRCR00", "José Luis", "García Hernández", date(1990, 1, 1)
        )
        self.assertTrue(any("María/José" in e for e in errs))


class SuspensionJugadorTest(TestCase):
    def setUp(self):
        self.cat = Categoria.objects.create(
            nombre="Libre", dias_juego=["SAB"], curp_obligatoria=False
        )
        self.campo = Campo.objects.create(nombre="Campo 1", activo=True)
        self.eq1 = Equipo.objects.create(nombre="Aguilas", categoria=self.cat, activo=True)
        self.eq2 = Equipo.objects.create(nombre="Toros", categoria=self.cat, activo=True)
        self.jugador = Jugador.objects.create(
            nombre="Juan", apellido="Perez", equipo=self.eq1, posicion="DEL"
        )
        self.temporada = Temporada.objects.create(
            categoria=self.cat, nombre="Temp 1", fecha_inicio=date.today(),
            tipo_rol="TODOS", vueltas=1, iniciada=True,
        )

    def _jornada(self, numero):
        return Jornada.objects.create(temporada=self.temporada, numero=numero, nombre=f"J{numero}")

    def _partido(self, jornada, local, visit, fecha, n=1):
        return Partido.objects.create(
            temporada=self.temporada, jornada=jornada, equipo_local=local,
            equipo_visitante=visit, fecha_hora=fecha, campo=self.campo, estado="FIN",
        )

    def test_restantes_cuenta_partidos_finalizados(self):
        susp = SuspensionJugador.objects.create(
            jugador=self.jugador, categoria=self.cat, equipo=self.eq1,
            temporada=self.temporada, jornadas=2, fecha_inicio=date.today(),
        )
        self.assertEqual(susp.restantes(), 2)
        j = self._jornada(1)
        self._partido(j, self.eq1, self.eq2, timezone.now(), n=1)
        self.assertEqual(susp.restantes(), 1)
        j2 = self._jornada(2)
        self._partido(j2, self.eq2, self.eq1, timezone.now() + timedelta(days=7), n=2)
        self.assertEqual(susp.restantes(), 0)
        self.assertFalse(susp.vigente())

    def test_tabla_castigados_incluye_suspension_manual(self):
        SuspensionJugador.objects.create(
            jugador=self.jugador, categoria=self.cat, equipo=self.eq1,
            temporada=self.temporada, jornadas=2, fecha_inicio=date.today(),
        )
        client = Client()
        r = client.get(f"/tabla-castigados/?temporada={self.temporada.id}")
        self.assertEqual(r.status_code, 200)
        # La página usa el patrón sin permiso (ruta pública de tabla)
        self.assertContains(r, "Juan Perez")

    def test_tabla_castigados_heredero_global_sin_temporada(self):
        # Heredado global (equipo None) en una categoría sin temporadas:
        # debe verse igual aunque no haya temporada creada
        cat2 = Categoria.objects.create(nombre="Libre B", curp_obligatoria=False)
        SuspensionJugador.objects.create(
            jugador=self.jugador, categoria=cat2, equipo=None,
            jornadas=2, fecha_inicio=date.today(),
        )
        r = Client().get(f"/tabla-castigados/?categoria={cat2.id}")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Juan Perez")
        self.assertContains(r, "Expulsado de la liga")

    def test_exportar_castigados_con_suspension_sin_equipo(self):
        # Suspensión manual sin equipo (expulsado de la liga): PDF y Excel no deben fallar
        SuspensionJugador.objects.create(
            jugador=self.jugador, categoria=self.cat, equipo=None,
            temporada=self.temporada, jornadas=2, fecha_inicio=date.today(),
        )
        from django.contrib.auth import get_user_model
        User = get_user_model()
        User.objects.create_superuser(username="adminc", password="p")
        client = Client()
        client.force_login(User.objects.get(username="adminc"))
        for fmt in ("pdf", "xlsx"):
            with self.subTest(fmt=fmt):
                r = client.get(f"/reportes/castigados/{fmt}/?temporada={self.temporada.id}")
                self.assertEqual(r.status_code, 200)
                self.assertGreater(len(r.content), 1000)

    def test_exportar_castigados_con_suspension_vitalicia_sin_equipo(self):
        SuspensionJugador.objects.create(
            jugador=self.jugador, categoria=self.cat, equipo=None,
            vitalicia=True, fecha_inicio=date.today(),
        )
        from django.contrib.auth import get_user_model
        User = get_user_model()
        User.objects.create_superuser(username="admincv", password="p")
        client = Client()
        client.force_login(User.objects.get(username="admincv"))
        for fmt in ("pdf", "xlsx"):
            with self.subTest(fmt=fmt):
                r = client.get(f"/reportes/castigados/{fmt}/?temporada={self.temporada.id}")
                self.assertEqual(r.status_code, 200)
                self.assertGreater(len(r.content), 1000)

    def test_tabla_castigados_usa_categoria_preferida_del_usuario(self):
        from django.contrib.auth import get_user_model
        cat_principal = Categoria.objects.create(
            nombre="Principal", activo=True, es_principal=True, curp_obligatoria=False
        )
        cat_pref = Categoria.objects.create(
            nombre="Preferida", activo=True, curp_obligatoria=False
        )
        Temporada.objects.create(
            categoria=cat_pref, nombre="T Pref", fecha_inicio=date.today(),
            tipo_rol="TODOS", vueltas=1,
        )
        SuspensionJugador.objects.create(
            jugador=self.jugador, categoria=cat_pref, equipo=None,
            jornadas=2, fecha_inicio=date.today(),
        )
        User = get_user_model()
        u = User.objects.create_user(
            username="prefcat", password="x", categoria_preferida=cat_pref
        )
        client = Client()
        client.force_login(u)
        r = client.get("/tabla-castigados/")
        self.assertEqual(r.status_code, 200)
        # Sin seleccionar categoría se usa la preferida del usuario, no la principal
        self.assertContains(r, "Juan Perez")

    def test_clean_equipo_bloquea_cambio_de_equipo_suspendido(self):
        from .forms import JugadorForm
        SuspensionJugador.objects.create(
            jugador=self.jugador, categoria=self.cat, equipo=self.eq1,
            temporada=self.temporada, jornadas=1, fecha_inicio=date.today(),
        )
        form = JugadorForm(
            data={
                "nombre": "Juan", "apellido": "Perez", "posicion": "DEL",
                "equipo": self.eq2.id, "tipo_documento": "CURP",
            },
            instance=self.jugador,
        )
        self.assertFalse(form.is_valid())
        self.assertIn("suspensión activa", str(form.errors.get("equipo", "")))

    def test_levantar_suspension_desactiva(self):
        susp = SuspensionJugador.objects.create(
            jugador=self.jugador, categoria=self.cat, equipo=self.eq1,
            temporada=self.temporada, jornadas=3, fecha_inicio=date.today(),
        )
        susp.activo = False
        susp.save(update_fields=["activo"])
        self.assertEqual(susp.restantes(), 0)


class MovimientosEquipoTest(TestCase):
    def setUp(self):
        self.c1 = Categoria.objects.create(nombre="Primera", nivel=1)
        self.c2 = Categoria.objects.create(nombre="Intermedia", nivel=2)
        self.c3 = Categoria.objects.create(nombre="Segunda", nivel=3)
        self.eq1 = Equipo.objects.create(nombre="T1", categoria=self.c1, activo=True)
        self.eq2 = Equipo.objects.create(nombre="T2", categoria=self.c2, activo=True)
        self.j = Jugador.objects.create(nombre="Juan", apellido="Perez", equipo=self.eq1, fecha_nacimiento=date(1990, 1, 1))

    def _temporada(self, **kw):
        kw.setdefault("fecha_inicio", date(2026, 6, 1))
        kw.setdefault("tipo_rol", "TODOS")
        kw.setdefault("aplicar_movimientos", True)
        return Temporada.objects.create(categoria=self.c1, nombre="Temp 2026", **kw)

    def test_descenso_restringe_categoria_superior(self):
        mov = MovimientoEquipo.objects.create(
            temporada=self._temporada(), equipo=self.eq1,
            tipo="DESCENSO", origen_categoria=self.c1, destino_categoria=self.c2,
        )
        errs_sup = errores_movimiento_jugador(self.j, self.eq2, self.c1)
        self.assertTrue(any("descendió" in e for e in errs_sup))
        self.assertEqual(errores_movimiento_jugador(self.j, self.eq2, self.c2), [])
        self.assertEqual(errores_movimiento_jugador(self.j, self.eq2, self.c3), [])
        self.assertTrue(mov.cupo_50() == 0)

    def test_desaparicion_restringe_por_encima_y_dos_niveles_abajo(self):
        MovimientoEquipo.objects.create(
            temporada=self._temporada(), equipo=self.eq1,
            tipo="DESAPARECE", origen_categoria=self.c1,
        )
        self.assertTrue(any("desapareció" in e for e in errores_movimiento_jugador(self.j, None, self.c3)))
        self.assertEqual(errores_movimiento_jugador(self.j, None, self.c1), [])
        self.assertEqual(errores_movimiento_jugador(self.j, None, self.c2), [])

    def test_ascenso_50_por_ciento_fifo(self):
        c_sup = Categoria.objects.create(nombre="Maxima", nivel=0)
        eq_sup = Equipo.objects.create(nombre="T1S", categoria=c_sup, activo=True)
        otros = [
            Jugador.objects.create(nombre=f"P{i}", apellido="X", equipo=self.eq1, fecha_nacimiento=date(1990, 1, 1))
            for i in range(3)
        ]
        mov = MovimientoEquipo.objects.create(
            temporada=self._temporada(), equipo=self.eq1,
            tipo="ASCENSO", origen_categoria=self.c1, destino_categoria=c_sup,
            jugadores_plantilla=4,
        )
        # Cambiarse a OTRO equipo en categoria superior ok (cupo libre, categoría válida)
        j2 = otros[0]
        eq_destino = eq_sup
        self.assertEqual(errores_movimiento_jugador(j2, eq_destino, c_sup), [])
        # FIFO: registramos 2 (50% de 4) en el equipo de la categoría superior
        for o in otros[:2]:
            JugadorEquipo.objects.create(jugador=o, equipo=eq_sup, activo=True, es_principal=True)
        self.assertEqual(mov.transferidos(), 2)
        self.assertEqual(mov.cupo_50(), 2)
        errs = errores_movimiento_jugador(otros[2], eq_sup, c_sup)
        self.assertTrue(any("cupo del 50%" in e for e in errs))
        # Categoría incorrecta bloqueada aunque haya cupo
        errs_cat = errores_movimiento_jugador(self.j, self.eq2, self.c2)
        self.assertTrue(any("solo puede" in e for e in errs_cat))

    def test_se_queda_no_restringe(self):
        MovimientoEquipo.objects.create(
            temporada=self._temporada(), equipo=self.eq1,
            tipo="SE_QUEDA", origen_categoria=self.c1,
        )
        self.assertEqual(errores_movimiento_jugador(self.j, self.eq2, self.c3), [])
        self.assertEqual(errores_movimiento_jugador(self.j, self.eq2, self.c1), [])

    def test_guardar_movimientos_mueve_categoria_y_desactiva(self):
        from django.contrib.auth import get_user_model
        catA = Categoria.objects.create(nombre="Maxima", nivel=19)
        catB = Categoria.objects.create(nombre="Intermedia", nivel=20)
        catC = Categoria.objects.create(nombre="Segunda", nivel=21)
        t = Temporada.objects.create(
            categoria=catB, nombre="Temp", fecha_inicio=date(2026, 6, 1),
            tipo_rol="TODOS", num_ascensos=1, num_descensos=1, aplicar_movimientos=True,
            finalizada=True,
        )
        eqA = Equipo.objects.create(nombre="Alpha", categoria=catB, activo=True)
        eqB = Equipo.objects.create(nombre="Beta", categoria=catB, activo=True)
        Partido.objects.create(
            temporada=t, equipo_local=eqA, equipo_visitante=eqB,
            fecha_hora=timezone.make_aware(datetime(2026, 6, 5, 12, 0)),
            estado="FIN", goles_local=2, goles_visitante=1,
        )
        User = get_user_model()
        admin = User.objects.create_superuser(username="admin", password="p")
        self.client.force_login(admin)
        r = self.client.get(f"/temporadas/{t.pk}/movimientos/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Sube a Maxima")
        r_post = self.client.post(f"/temporadas/{t.pk}/movimientos/", {
            f"tipo_{eqA.pk}": "ASCENSO", f"tipo_{eqB.pk}": "DESAPARECE",
        })
        self.assertEqual(r_post.status_code, 302)
        eqA.refresh_from_db(); eqB.refresh_from_db()
        self.assertEqual(eqA.categoria_id, catA.id)  # sube a Maxima
        self.assertFalse(eqB.activo)  # desaparece
        movA = MovimientoEquipo.objects.get(equipo=eqA)
        self.assertEqual(movA.tipo, "ASCENSO")
        self.assertEqual(movA.destino_categoria_id, catA.id)
        movB = MovimientoEquipo.objects.get(equipo=eqB)
        self.assertEqual(movB.tipo, "DESAPARECE")
        self.assertEqual(movB.origen_categoria_id, catB.id)

    def test_regla_inactiva_suspende_restriccion(self):
        MovimientoEquipo.objects.create(
            temporada=self._temporada(), equipo=self.eq1,
            tipo="DESCENSO", origen_categoria=self.c1, destino_categoria=self.c2,
            regla_activa=False,
        )
        self.assertEqual(errores_movimiento_jugador(self.j, self.eq2, self.c1), [])
        mov = MovimientoEquipo.objects.get(equipo=self.eq1)
        mov.regla_activa = True
        mov.save(update_fields=["regla_activa"])
        self.assertTrue(any("descendió" in e for e in errores_movimiento_jugador(self.j, self.eq2, self.c1)))

    def test_cupo_porcentaje_configurable(self):
        mov = MovimientoEquipo.objects.create(
            temporada=self._temporada(), equipo=self.eq1,
            tipo="ASCENSO", origen_categoria=self.c1, destino_categoria=self.c2,
            jugadores_plantilla=4, cupo_porcentaje=25,
        )
        self.assertEqual(mov.cupo_50(), 1)
        mov.cupo_porcentaje = 100
        mov.save(update_fields=["cupo_porcentaje"])
        self.assertEqual(mov.cupo_50(), 4)

    def test_movimientos_guardan_vigencia_y_cupo(self):
        from django.contrib.auth import get_user_model
        catA = Categoria.objects.create(nombre="MaximaV", nivel=31)
        catB = Categoria.objects.create(nombre="IntermediaV", nivel=32)
        Categoria.objects.create(nombre="SegundaV", nivel=33)
        t = Temporada.objects.create(
            categoria=catB, nombre="Temp V", fecha_inicio=date(2026, 6, 1),
            tipo_rol="TODOS", aplicar_movimientos=True,
        )
        eqA = Equipo.objects.create(nombre="AlphaV", categoria=catB, activo=True)
        eqB = Equipo.objects.create(nombre="BetaV", categoria=catB, activo=True)
        Partido.objects.create(
            temporada=t, equipo_local=eqA, equipo_visitante=eqB,
            fecha_hora=timezone.make_aware(datetime(2026, 6, 5, 12, 0)),
            estado="FIN", goles_local=1, goles_visitante=0,
        )
        User = get_user_model()
        admin = User.objects.create_superuser(username="adminv", password="p")
        self.client.force_login(admin)
        r = self.client.post(f"/temporadas/{t.pk}/movimientos/", {
            f"tipo_{eqA.pk}": "ASCENSO",
            f"regla_activa_{eqA.pk}": "1",
            f"cupo_pct_{eqA.pk}": "25",
        })
        self.assertEqual(r.status_code, 302)
        mov = MovimientoEquipo.objects.get(temporada=t, equipo=eqA)
        self.assertTrue(mov.regla_activa)
        self.assertEqual(mov.cupo_porcentaje, 25)
        # Segunda pasada sin marcar la regla: se apaga la vigencia
        r2 = self.client.post(f"/temporadas/{t.pk}/movimientos/", {
            f"tipo_{eqA.pk}": "ASCENSO",
            f"cupo_pct_{eqA.pk}": "25",
        })
        self.assertEqual(r2.status_code, 302)
        mov.refresh_from_db()
        self.assertFalse(mov.regla_activa)


class JugadorSinEquipoTest(TestCase):
    def setUp(self):
        self.cat = Categoria.objects.create(
            nombre="Libre", dias_juego=["SAB"], curp_obligatoria=False
        )
        self.campo = Campo.objects.create(nombre="Campo 1", activo=True)
        self.eq1 = Equipo.objects.create(nombre="Aguilas", categoria=self.cat, activo=True)
        self.eq2 = Equipo.objects.create(nombre="Toros", categoria=self.cat, activo=True)
        self.temporada = Temporada.objects.create(
            categoria=self.cat, nombre="Temp 1", fecha_inicio=date.today(),
            tipo_rol="TODOS", vueltas=1, iniciada=True,
        )

    def _jornada(self, numero):
        return Jornada.objects.create(temporada=self.temporada, numero=numero, nombre=f"J{numero}")

    def _partido(self, jornada, local, visit, fecha, n=1):
        return Partido.objects.create(
            temporada=self.temporada, jornada=jornada, equipo_local=local,
            equipo_visitante=visit, fecha_hora=fecha, campo=self.campo, estado="FIN",
        )

    def test_jugador_sin_equipo_y_asignacion_posterior(self):
        j = Jugador.objects.create(nombre="Ana", apellido="Lopez", posicion="DEL")
        self.assertIsNone(j.equipo)
        self.assertFalse(j.registros_equipo.exists())
        j.equipo = self.eq1
        j.save(update_fields=["equipo"])
        self.assertTrue(j.registros_equipo.filter(equipo=self.eq1, es_principal=True).exists())

    def test_suspension_sin_equipo_cuenta_jornadas_de_la_categoria(self):
        j = Jugador.objects.create(nombre="Beto", apellido="Sanz", posicion="DEL")
        susp = SuspensionJugador.objects.create(
            jugador=j, categoria=self.cat, equipo=None,
            temporada=self.temporada, jornadas=2, fecha_inicio=date.today(),
        )
        self.assertEqual(susp.restantes(), 2)
        j1 = self._jornada(1)
        # Dos partidos FIN en la misma jornada descontan una sola jornada
        self._partido(j1, self.eq1, self.eq2, timezone.now(), n=1)
        self._partido(j1, self.eq2, self.eq1, timezone.now() + timedelta(hours=1), n=2)
        self.assertEqual(susp.consumidos(), 1)
        self.assertEqual(susp.restantes(), 1)
        j2 = self._jornada(2)
        self._partido(j2, self.eq1, self.eq2, timezone.now() + timedelta(days=7), n=3)
        self.assertEqual(susp.restantes(), 0)


class JugadorHeredadoTest(TestCase):
    def setUp(self):
        self.c1 = Categoria.objects.create(nombre="MaximaH", nivel=0, activo=True, curp_obligatoria=False)
        self.c2 = Categoria.objects.create(nombre="IntermediaH", nivel=1, activo=True, curp_obligatoria=False)
        self.c3 = Categoria.objects.create(nombre="SegundaH", nivel=2, activo=True, curp_obligatoria=False)
        self.eq_sup = Equipo.objects.create(nombre="Kings", categoria=self.c1, activo=True)

    def test_castigado_con_jornadas_crea_suspension_global(self):
        j = Jugador.objects.create(nombre="Leo", apellido="Messi", posicion="DEL")
        he = aplicar_movimiento_a_jugador(j, "CASTIGADO", self.c2, jornadas=2, motivo="Expulsado")
        self.assertEqual(JugadorHerencia.objects.filter(jugador=j, activo=True).count(), 1)
        susp = SuspensionJugador.objects.filter(jugador=j, activo=True).first()
        self.assertIsNotNone(susp)
        self.assertEqual(susp.equipo, None)
        self.assertEqual(susp.jornadas, 2)
        self.assertEqual(he.tipo, "CASTIGADO")

    def test_castigado_sin_jornadas_no_crea_suspension(self):
        j = Jugador.objects.create(nombre="Sergio", apellido="Ramos", posicion="DEF")
        aplicar_movimiento_a_jugador(j, "CASTIGADO", self.c2, jornadas=0)
        self.assertFalse(SuspensionJugador.objects.filter(jugador=j, activo=True).exists())
        self.assertEqual(errores_movimiento_jugador(j, self.eq_sup, self.c1), [])

    def test_expulsado_por_vida_crea_suspension_permanente(self):
        j = Jugador.objects.create(nombre="Diego", apellido="Campos Cerrito", posicion="DEL")
        he = aplicar_movimiento_a_jugador(j, "VITALICIO", self.c2, motivo="Expulsado por agresión")
        self.assertEqual(he.tipo, "VITALICIO")
        susp = SuspensionJugador.objects.filter(jugador=j, activo=True).first()
        self.assertIsNotNone(susp)
        self.assertTrue(susp.vitalicia)
        self.assertEqual(susp.equipo, None)
        # Nunca se cumple: restantes siempre > 0 aunque pasen partidos
        self.assertEqual(susp.restantes(), 1)
        self.assertTrue(susp.vigente())
        # Bloquea el registro en cualquier categoría
        errs = errores_movimiento_jugador(j, self.eq_sup, self.c1)
        self.assertTrue(any("por vida" in e for e in errs))
        errs2 = errores_movimiento_jugador(j, None, self.c3)
        self.assertTrue(any("por vida" in e for e in errs2))

    def test_expulsado_por_vida_sale_de_todos_los_equipos(self):
        j = Jugador.objects.create(
            nombre="Raul", apellido="Jimeno", posicion="DEF", equipo=self.eq_sup
        )
        aplicar_movimiento_a_jugador(j, "VITALICIO", self.c1)
        j.refresh_from_db()
        self.assertIsNone(j.equipo)
        self.assertFalse(JugadorEquipo.objects.filter(jugador=j, activo=True).exists())
        self.assertEqual(SuspensionJugador.objects.filter(jugador=j, activo=True).count(), 1)

    def test_castigado_heredero_expulsa_de_todos_los_equipos(self):
        j = Jugador.objects.create(
            nombre="Gabriel", apellido="Reyes Casas", posicion="DEL", equipo=self.eq_sup
        )
        otro = Equipo.objects.create(nombre="Otro Club", categoria=self.c1, activo=True)
        JugadorEquipo.objects.create(jugador=j, equipo=otro, es_principal=False, activo=True)
        aplicar_movimiento_a_jugador(j, "CASTIGADO", self.c1, jornadas=2, motivo="Expulsado")
        j.refresh_from_db()
        self.assertIsNone(j.equipo)
        # Ningún registro activo: no aparece en rosters ni credenciales
        self.assertFalse(JugadorEquipo.objects.filter(jugador=j, activo=True).exists())
        # Corte histórico conservado
        self.assertEqual(JugadorEquipo.objects.filter(jugador=j, equipo=self.eq_sup).count(), 1)
        self.assertEqual(SuspensionJugador.objects.filter(jugador=j, activo=True).count(), 1)

    def test_herencia_ascenso_descenso_desaparece_aplican_reglas(self):
        base = Jugador.objects.create(nombre="Pelusa", apellido="Rojo", posicion="DEL")
        aplicar_movimiento_a_jugador(base, "ASCENSO", self.c3)
        # Ascendió a SegundaH (c3): solo puede jugar en c3 o la inmediata inferior
        eq_c3 = Equipo.objects.create(nombre="SegundaH Eq", categoria=self.c3, activo=True)
        self.assertEqual(errores_movimiento_jugador(base, eq_c3, self.c3), [])
        self.assertTrue(any("asciende (heredado)" in e for e in errores_movimiento_jugador(base, self.eq_sup, self.c1)))
        self.assertTrue(any("asciende (heredado)" in e for e in errores_movimiento_jugador(base, self.eq_sup, self.c2)))

        des = Jugador.objects.create(nombre="Dena", apellido="Dos", posicion="MED")
        aplicar_movimiento_a_jugador(des, "DESCENSO", self.c1)
        # Descendió a MaximaH (c1): no puede jugar por encima (no hay); sí en c1, c2, c3
        self.assertEqual(errores_movimiento_jugador(des, None, self.c1), [])
        eq_c2 = Equipo.objects.create(nombre="IntermediaH Eq", categoria=self.c2, activo=True)
        self.assertEqual(errores_movimiento_jugador(des, eq_c2, self.c2), [])
        self.assertEqual(errores_movimiento_jugador(des, None, self.c3), [])

        desa = Jugador.objects.create(nombre="Baja", apellido="Tres", posicion="POR")
        aplicar_movimiento_a_jugador(desa, "DESAPARECE", self.c1)
        self.assertEqual(errores_movimiento_jugador(desa, None, self.c1), [])
        self.assertEqual(errores_movimiento_jugador(desa, eq_c2, self.c2), [])
        self.assertTrue(any("equipo dado de baja" in e for e in errores_movimiento_jugador(desa, None, self.c3)))

    def test_herencia_desactivada_deja_de_restringir(self):
        j = Jugador.objects.create(nombre="Free", apellido="Cuatro", posicion="DEL")
        he = aplicar_movimiento_a_jugador(j, "DESAPARECE", self.c1)
        self.assertTrue(any("equipo dado de baja" in e for e in errores_movimiento_jugador(j, None, self.c3)))
        he.activo = False
        he.save(update_fields=["activo"])
        self.assertEqual(errores_movimiento_jugador(j, None, self.c3), [])

    def test_alta_duplicado_bloqueado_y_guardado_herencia(self):
        from django.contrib.auth import get_user_model
        existente = Jugador.objects.create(
            nombre="Dup", apellido="Cinco", posicion="DEL", equipo=self.eq_sup,
            fecha_nacimiento=date(1995, 5, 5),
        )
        User = get_user_model()
        admin = User.objects.create_superuser(username="adminh", password="p")
        self.client.force_login(admin)
        # Crear con mismo nombre+fecha: debe bloquear el alta (responder 200 sin crear otro)
        r = self.client.post("/jugadores-heredados/", {
            "accion": "crear",
            "nombre": "Dup", "apellido": "Cinco",
            "fecha_nacimiento": "1995-05-05",
            "tipo": "CASTIGADO", "categoria_id": self.c1.id, "jornadas": "1",
        })
        self.assertEqual(r.status_code, 200)
        self.assertEqual(Jugador.objects.filter(nombre="Dup").count(), 1)
        self.assertContains(r, "Ya existe")
        # Heredar al existente sin nuevo alta
        r2 = self.client.post("/jugadores-heredados/", {
            "accion": "heredar", "jugador_id": existente.id,
            "tipo": "CASTIGADO", "categoria_id": self.c1.id, "jornadas": "1", "motivo": "",
        })
        self.assertEqual(r2.status_code, 302)
        self.assertTrue(SuspensionJugador.objects.filter(jugador=existente, activo=True).exists())
        # Expulsado por castigo heredado: sale del equipo actual y deja el cupo
        existente.refresh_from_db()
        self.assertIsNone(existente.equipo)
        self.assertFalse(
            JugadorEquipo.objects.filter(jugador=existente, equipo=self.eq_sup, activo=True).exists()
        )

    def test_crear_jugador_nuevo_con_herencia(self):
        from django.contrib.auth import get_user_model
        User = get_user_model()
        admin = User.objects.create_superuser(username="adminj", password="p")
        self.client.force_login(admin)
        r = self.client.post("/jugadores-heredados/", {
            "accion": "crear",
            "nombre": "Nuevo", "apellido": "Heredado",
            "tipo": "DESCENSO", "categoria_id": self.c2.id, "jornadas": "0", "motivo": "baja previa",
        })
        self.assertEqual(r.status_code, 302)
        j = Jugador.objects.get(nombre="Nuevo")
        self.assertIsNone(j.equipo)
        he = JugadorHerencia.objects.get(jugador=j)
        self.assertEqual(he.tipo, "DESCENSO")
        self.assertEqual(he.categoria_id, self.c2.id)

    def test_busqueda_nombre_completo_con_acentos_sin_acento(self):
        from django.contrib.auth import get_user_model
        User = get_user_model()
        admin = User.objects.create_superuser(username="adminbus", password="p")
        self.client.force_login(admin)
        j = Jugador.objects.create(
            nombre="Juan Carlos", apellido="Pérez Ramírez", posicion="DEL",
            equipo=self.eq_sup, fecha_nacimiento=date(1990, 1, 1),
        )
        # Nombre completo en otro orden y sin tildes debe encontrarlo
        r = self.client.get("/jugadores-heredados/", {"q": "juan carlos perez ramirez"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Pérez Ramírez")
        # Buscando solo el apellido compuesto también
        r2 = self.client.get("/jugadores-heredados/", {"q": "perez ramirez"})
        self.assertContains(r2, j.apellido)

    def test_duplicado_solo_por_nombre_y_apellido_bloqueado(self):
        from django.contrib.auth import get_user_model
        User = get_user_model()
        admin = User.objects.create_superuser(username="admindup2", password="p")
        self.client.force_login(admin)
        existente = Jugador.objects.create(nombre="Ana", apellido="López", posicion="DEL", equipo=self.eq_sup)
        # Sin CURP ni fecha: solo nombre+apellido (con tilde distinta) debe bloquear
        r = self.client.post("/jugadores-heredados/", {
            "accion": "crear",
            "nombre": "Ana", "apellido": "Lopez",
            "tipo": "CASTIGADO", "categoria_id": self.c1.id, "jornadas": "0",
        })
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Ya existe")
        self.assertEqual(Jugador.objects.filter(nombre="Ana").count(), 1)

    def test_crear_forzado_permite_homonimo(self):
        from django.contrib.auth import get_user_model
        User = get_user_model()
        admin = User.objects.create_superuser(username="adminfor", password="p")
        self.client.force_login(admin)
        Jugador.objects.create(nombre="Ana", apellido="López", posicion="DEL", equipo=self.eq_sup)
        r = self.client.post("/jugadores-heredados/", {
            "accion": "crear", "forzar": "1",
            "nombre": "Ana", "apellido": "Lopez",
            "tipo": "CASTIGADO", "categoria_id": self.c1.id, "jornadas": "0", "motivo": "homonimo real",
        })
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Jugador.objects.filter(nombre="Ana").count(), 2)

    def test_heredar_corrige_reemplazando_suspension_anterior(self):
        from django.contrib.auth import get_user_model
        User = get_user_model()
        admin = User.objects.create_superuser(username="admincor", password="p")
        self.client.force_login(admin)
        j = Jugador.objects.create(nombre="Alexis", apellido="Montoya", posicion="DEL")
        # Primera herencia: castigado en c2
        self.client.post("/jugadores-heredados/", {
            "accion": "heredar", "jugador_id": j.id,
            "tipo": "CASTIGADO", "categoria_id": self.c2.id, "jornadas": "2", "motivo": "",
        })
        self.assertTrue(SuspensionJugador.objects.filter(jugador=j, activo=True).exists())
        # Corregir: heredar de nuevo con otra categoría no debe bloquear
        r = self.client.post("/jugadores-heredados/", {
            "accion": "heredar", "jugador_id": j.id,
            "tipo": "CASTIGADO", "categoria_id": self.c1.id, "jornadas": "3", "motivo": "ajustado",
        })
        self.assertEqual(r.status_code, 302)
        # Reemplazo: queda una sola herencia activa y un solo registro en total
        self.assertEqual(SuspensionJugador.objects.filter(jugador=j, activo=True).count(), 1)
        susp = SuspensionJugador.objects.get(jugador=j, activo=True)
        self.assertEqual(susp.categoria_id, self.c1.id)
        self.assertEqual(susp.jornadas, 3)
        self.assertEqual(JugadorHerencia.objects.filter(jugador=j).count(), 1)
        he_new = JugadorHerencia.objects.get(jugador=j)
        self.assertTrue(he_new.activo)
        self.assertEqual(he_new.categoria_id, self.c1.id)

    def test_eliminar_castigado_levanta_la_suspension(self):
        from django.contrib.auth import get_user_model
        User = get_user_model()
        admin = User.objects.create_superuser(username="admindes", password="p")
        self.client.force_login(admin)
        j = Jugador.objects.create(nombre="Beto", apellido="Des", posicion="MED")
        self.client.post("/jugadores-heredados/", {
            "accion": "heredar", "jugador_id": j.id,
            "tipo": "CASTIGADO", "categoria_id": self.c2.id, "jornadas": "2", "motivo": "",
        })
        self.assertTrue(SuspensionJugador.objects.filter(jugador=j, activo=True).exists())
        he = JugadorHerencia.objects.get(jugador=j)
        r = self.client.post("/jugadores-heredados/", {
            "accion": "eliminar", "herencia_id": he.id,
        })
        self.assertEqual(r.status_code, 302)
        self.assertFalse(JugadorHerencia.objects.filter(jugador=j).exists())
        self.assertFalse(SuspensionJugador.objects.filter(jugador=j, activo=True).exists())

    def test_herencia_desaparece_saca_del_equipo_incompatible(self):
        from django.contrib.auth import get_user_model
        User = get_user_model()
        admin = User.objects.create_superuser(username="adminmov", password="p")
        self.client.force_login(admin)
        # Orden: c1=Primera(nivel 0), c2=Intermedia(1), c3=Segunda(2)
        eq_3 = Equipo.objects.create(nombre="SegundaH Club", categoria=self.c3, activo=True)
        j = Jugador.objects.create(nombre="Dani", apellido="Oli", posicion="DEL", equipo=eq_3)
        rer = self.client.post("/jugadores-heredados/", {
            "accion": "heredar", "jugador_id": j.id,
            "tipo": "DESAPARECE", "categoria_id": self.c1.id, "jornadas": "0", "motivo": "",
        })
        self.assertEqual(rer.status_code, 302)
        j.refresh_from_db()
        self.assertIsNone(j.equipo)
        # Registro desactivado (historial conservado) y no aparece en listas
        self.assertFalse(JugadorEquipo.objects.filter(jugador=j, equipo=eq_3, activo=True).exists())
        self.assertEqual(JugadorEquipo.objects.filter(jugador=j, equipo=eq_3).count(), 1)

    def test_herencia_descenso_conserva_equipo_compatible(self):
        from django.contrib.auth import get_user_model
        User = get_user_model()
        admin = User.objects.create_superuser(username="admincon", password="p")
        self.client.force_login(admin)
        eq_3 = Equipo.objects.create(nombre="SegundaH Club B", categoria=self.c3, activo=True)
        j = Jugador.objects.create(nombre="Ana", apellido="Com", posicion="DEF", equipo=eq_3)
        r = self.client.post("/jugadores-heredados/", {
            "accion": "heredar", "jugador_id": j.id,
            "tipo": "DESCENSO", "categoria_id": self.c1.id, "jornadas": "0", "motivo": "",
        })
        self.assertEqual(r.status_code, 302)
        j.refresh_from_db()
        # c3 está por debajo (nivel 2): descendido en c1(nivel 0) sí puede jugar ahí
        self.assertEqual(j.equipo_id, eq_3.id)
        self.assertTrue(JugadorEquipo.objects.filter(jugador=j, equipo=eq_3, activo=True).exists())


class FotoJugadorTest(TestCase):
    """Vista subir_foto_jugador: solo permite subir a jugadores SIN foto.
    El reemplazo de una foto existente solo lo hace superusuario o con
    período de altas activo."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        import tempfile, shutil
        from django.core.files.storage import FileSystemStorage, default_storage
        cls._shutil = shutil
        cls._tmp_media = tempfile.mkdtemp()
        cls._prev_storage = default_storage._wrapped
        default_storage._wrapped = FileSystemStorage(location=cls._tmp_media, base_url="/media/")

    @classmethod
    def tearDownClass(cls):
        from django.core.files.storage import default_storage
        default_storage._wrapped = cls._prev_storage
        cls._shutil.rmtree(cls._tmp_media, ignore_errors=True)
        super().tearDownClass()

    def _png(self, nombre):
        from django.core.files.uploadedfile import SimpleUploadedFile
        import io
        from PIL import Image
        buf = io.BytesIO()
        Image.new("RGB", (8, 8), "red").save(buf, format="PNG")
        buf.seek(0)
        return SimpleUploadedFile(nombre, buf.getvalue(), content_type="image/png")

    def setUp(self):
        from django.contrib.auth import get_user_model
        from accounts.models import Rol
        self._shutil.rmtree(self._tmp_media, ignore_errors=True)
        self.rol = Rol.objects.create(
            nombre="CapturaFoto", permisos={"gestion_jugadores": True, "jugador_foto": True}
        )
        self.user = get_user_model().objects.create_user(
            username="capturafoto", password="p", rol=self.rol
        )
        self.cat = Categoria.objects.create(nombre="Libre", dias_juego=["SAB"])
        self.eq = Equipo.objects.create(nombre="Aguilas", categoria=self.cat, activo=True)
        self.j_sin_foto = Jugador.objects.create(nombre="Juan", apellido="Perez", equipo=self.eq)
        self.j_con_foto = Jugador.objects.create(nombre="Ana", apellido="Lopez", equipo=self.eq)
        self.j_con_foto.foto = self._png("ana.png")
        self.j_con_foto.save()
        self.foto_original = self.j_con_foto.foto.name
        self.client.force_login(self.user)

    def test_sin_permiso_envia_403(self):
        from django.contrib.auth import get_user_model
        otro = get_user_model().objects.create_user(username="sinfoto", password="p")
        c = Client(raise_request_exception=False)
        c.force_login(otro)
        r = c.post(f"/jugadores/{self.j_sin_foto.pk}/subir-foto/", {"foto": self._png("x.png")})
        self.assertEqual(r.status_code, 403)

    def test_sube_foto_a_jugador_sin_foto(self):
        r = self.client.post(
            f"/jugadores/{self.j_sin_foto.pk}/subir-foto/", {"foto": self._png("juan.png")}
        )
        self.assertRedirects(r, "/jugadores/")
        self.j_sin_foto.refresh_from_db()
        self.assertTrue(self.j_sin_foto.foto)

    def test_no_reemplaza_foto_existente(self):
        r = self.client.post(
            f"/jugadores/{self.j_con_foto.pk}/subir-foto/", {"foto": self._png("nueva.png")}
        )
        self.assertRedirects(r, "/jugadores/")
        self.j_con_foto.refresh_from_db()
        self.assertEqual(self.j_con_foto.foto.name, self.foto_original)

    def test_superuser_si_reemplaza(self):
        from django.contrib.auth import get_user_model
        admin = get_user_model().objects.create_superuser(username="adminfoto", password="p")
        self.client.force_login(admin)
        r = self.client.post(
            f"/jugadores/{self.j_con_foto.pk}/subir-foto/", {"foto": self._png("nueva.png")}
        )
        self.assertRedirects(r, "/jugadores/")
        self.j_con_foto.refresh_from_db()
        self.assertTrue(self.j_con_foto.foto)
        self.assertNotEqual(self.j_con_foto.foto.name, self.foto_original)

    def test_rechaza_archivo_no_imagen(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        bad = SimpleUploadedFile("txt.txt", b"no soy imagen", content_type="text/plain")
        r = self.client.post(f"/jugadores/{self.j_sin_foto.pk}/subir-foto/", {"foto": bad})
        self.assertRedirects(r, "/jugadores/")
        self.j_sin_foto.refresh_from_db()
        self.assertFalse(self.j_sin_foto.foto)


class ImagenRolDashboardTest(TestCase):
    """El generador de imagen 'Próximos partidos' (todas las categorías)
    produce un PNG válido incluso sin logos."""

    def _secciones(self):
        return [{
            "categoria": "Primera",
            "temporada": "Torneo 2026",
            "jornada": "Jornada 3",
            "partidos": [{
                "local": "Equipo Local",
                "visitante": "Equipo Visita",
                "campo": "Cancha Norte",
                "fecha": "15/09 18:00",
                "logo_local": None,
                "logo_visitante": None,
            }],
            "descansan": ["Equipo Descansa"],
        }]

    def test_genera_png_valido(self):
        from league.social_image import generar_imagen_rol_dashboard
        buf = generar_imagen_rol_dashboard(self._secciones(), "11/09/2026 12:00")
        data = buf.getvalue()
        self.assertTrue(data.startswith(b"\x89PNG"))
        self.assertGreater(len(data), 1000)

    def test_sin_secciones_produce_imagen(self):
        from league.social_image import generar_imagen_rol_dashboard
        buf = generar_imagen_rol_dashboard([], "11/09/2026 12:00")
        self.assertTrue(buf.getvalue().startswith(b"\x89PNG"))


class ContadorVisitasTest(TestCase):
    """El contador de visitas agrupa la navegación de un mismo visitante
    dentro de una hora en UNA sola visita (inicio) y no cuenta al staff."""

    def test_visita_publica_se_registra(self):
        from .models import Visita
        antes = Visita.objects.count()
        r = self.client.get("/")
        self.assertEqual(r.status_code, 200)
        self.assertGreater(Visita.objects.count(), antes)
        self.assertGreater(Visita.objects.filter(path="/").count(), 0)
        self.assertTrue(Visita.objects.filter(path="/").latest("fecha").inicio)

    def test_staff_no_se_registra(self):
        from .models import Visita
        from django.contrib.auth import get_user_model
        staff = get_user_model().objects.create_user(username="staffvis", password="p", is_staff=True)
        self.client.force_login(staff)
        self.client.get("/")
        self.assertEqual(Visita.objects.count(), 0)

    def test_bot_no_se_registra(self):
        from .models import Visita
        self.client.get("/", HTTP_USER_AGENT="Googlebot/2.1 (+http://www.google.com/bot.html)")
        self.assertEqual(Visita.objects.count(), 0)

    def test_navegacion_en_misma_hora_es_una_sola_visita(self):
        from .models import Visita
        self.client.get("/")
        self.client.get("/tabla-castigados/")
        self.client.get("/")
        self.assertEqual(Visita.objects.filter(inicio=True).count(), 1)
        self.assertEqual(Visita.objects.count(), 3)

    def test_regreso_mas_de_una_hora_es_nueva_visita(self):
        from .models import Visita
        from django.utils import timezone
        from datetime import timedelta
        self.client.get("/")
        primera = Visita.objects.latest("fecha")
        primera.fecha = timezone.now() - timedelta(hours=2)
        primera.save(update_fields=["fecha"])
        self.client.get("/")
        # La nueva página es inicio (pasó más de 1 hora desde la última actividad)
        self.assertEqual(Visita.objects.filter(inicio=True).count(), 2)

    def test_contador_total_cuenta_inicios(self):
        from .models import Visita
        from django.contrib.auth import get_user_model
        self.client.get("/")
        self.client.get("/tabla-castigados/")
        Visita.objects.create(path="/", inicio=True)
        admin = get_user_model().objects.create_user(
            username="adminvis2", password="p", is_staff=True, is_superuser=True
        )
        self.client.force_login(admin)
        r = self.client.get("/push-logs/")
        self.assertEqual(r.status_code, 200)
        # 1 sesión real + 1 creada directamente
        self.assertEqual(r.context["visitas_unicas"], 2)
        self.assertEqual(r.context["paginas_mes"], 3)

    def test_pagina_push_logs_muestra_pestana_visitas(self):
        from django.contrib.auth import get_user_model
        from .models import Visita
        Visita.objects.create(path="/", ip="127.0.0.1", inicio=True)
        admin = get_user_model().objects.create_user(
            username="adminvis", password="p", is_staff=True, is_superuser=True
        )
        self.client.force_login(admin)
        r = self.client.get("/push-logs/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'id="visitas-tab"')
        self.assertContains(r, "Total de visitas")
        self.assertContains(r, "Páginas más visitadas")
        self.assertContains(r, "Inicio de visita")


"""Tests del feature 'agregar equipo secundario' desde /jugadores/.

Cubren dos cosas:
  1. Que el botón SOLO aparezca para jugadores que realmente pueden recibir un
     equipo secundario válido (categoría compatible + edad + período de altas).
  2. Que el endpoint aplique las mismas reglas y no toque datos del jugador.
"""
from datetime import date, timedelta

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from .models import Categoria, Equipo, Jugador, JugadorEquipo, SuspensionJugador, Temporada
from .views import _jugadores_con_secundario_posible


class EquipoSecundarioBaseTest(TestCase):
    def setUp(self):
        Usuario = get_user_model()

        # A (principal) ~ B (compatible con A). C no es compatible con nadie.
        self.cat_a = Categoria.objects.create(
            nombre="A", nivel=0, activo=True, curp_obligatoria=False,
            edad_minima=None, edad_maxima=None,
        )
        self.cat_b = Categoria.objects.create(
            nombre="B", nivel=1, activo=True, curp_obligatoria=False,
            edad_minima=None, edad_maxima=None,
        )
        self.cat_c = Categoria.objects.create(
            nombre="C", nivel=2, activo=True, curp_obligatoria=False,
            edad_minima=None, edad_maxima=None,
        )
        self.cat_a.categorias_compatibles.add(self.cat_b)
        self.cat_b.categorias_compatibles.add(self.cat_a)

        self.eq_a = Equipo.objects.create(nombre="EqA", categoria=self.cat_a, activo=True)
        self.eq_b = Equipo.objects.create(nombre="EqB", categoria=self.cat_b, activo=True)
        self.eq_c = Equipo.objects.create(nombre="EqC", categoria=self.cat_c, activo=True)

        self.jugador = Jugador.objects.create(
            nombre="Juan", apellido="Perez", posicion="DEL",
            equipo=self.eq_a, fecha_nacimiento=date(2000, 1, 1),
        )

        self.admin = Usuario.objects.create_user(
            username="adminsec", password="p", is_staff=True, is_superuser=True
        )
        self.url = reverse("jugador_agregar_equipo_secundario", args=[self.jugador.pk])

    def registrar(self, jugador, equipo, es_principal=False, activo=True):
        return JugadorEquipo.objects.create(
            jugador=jugador, equipo=equipo, es_principal=es_principal, activo=activo
        )


class BotonEquipoSecundarioTest(EquipoSecundarioBaseTest):
    """El botón se muestra únicamente si el jugador cumple los requisitos."""

    def test_se_muestra_si_hay_categoria_compatible(self):
        self.assertIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_se_muestra_si_ya_tiene_secundario_para_poder_quitarlo(self):
        # A solo es compatible con B y el jugador ya está en B: no hay categoría
        # nueva, pero el botón debe verse para poder darle de baja.
        self.registrar(self.jugador, self.eq_b)
        self.assertIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_no_se_muestra_si_no_hay_categoria_compatible_ni_secundarios(self):
        # Jugador en C, que no es compatible con nadie: sin rutas posibles.
        j = Jugador.objects.create(
            nombre="Solo", apellido="C", posicion="DEF",
            equipo=self.eq_c, fecha_nacimiento=date(2000, 1, 1),
        )
        self.assertNotIn(j.pk, _jugadores_con_secundario_posible([j]))

    def test_no_se_muestra_si_la_edad_no_cumple(self):
        self.cat_b.edad_minima = 60
        self.cat_b.save()
        # El jugador tiene ~25 años.
        self.assertNotIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_se_muestra_si_la_edad_cumple(self):
        self.cat_b.edad_minima = 18
        self.cat_b.edad_maxima = 40
        self.cat_b.save()
        self.assertIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_no_se_muestra_sin_equipo_principal(self):
        self.jugador.equipo = None
        self.jugador.save()
        self.assertNotIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_no_se_muestra_si_categoria_destino_no_tiene_equipos_activos(self):
        self.eq_b.activo = False
        self.eq_b.save()
        self.assertNotIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_no_se_muestra_si_no_hay_periodo_de_altas_en_temporada_en_curso(self):
        t = Temporada.objects.create(
            categoria=self.cat_b, nombre="T-B",
            fecha_inicio=date.today() - timedelta(days=30),
            iniciada=True, activa=True, finalizada=False,
        )
        # Temporada en curso SIN período de altas abierto -> no se puede dar de alta.
        self.assertNotIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))
        # Con período de altas abierto -> sí.
        from .models import PeriodoAltas
        PeriodoAltas.objects.create(
            temporada=t, tipo="fechas", activo=True,
            fecha_inicio=date.today(), fecha_fin=date.today() + timedelta(days=30)
        )
        self.assertIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_no_se_muestra_con_suspension_activa(self):
        SuspensionJugador.objects.create(
            jugador=self.jugador, categoria=self.cat_a, jornadas=2,
            fecha_inicio=date.today(),
        )
        self.assertNotIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_el_listado_solo_muestra_el_boton_a_quienes_cumplen(self):
        otro = Jugador.objects.create(
            nombre="Sin", apellido="Salida", posicion="DEF",
            equipo=self.eq_c, fecha_nacimiento=date(2000, 1, 1),
        )
        self.client.force_login(self.admin)
        r = self.client.get(reverse("jugador_list"))
        self.assertEqual(r.status_code, 200)
        puede = r.context["puede_secundario"]
        self.assertIn(self.jugador.pk, puede)
        self.assertNotIn(otro.pk, puede)


class EndpointEquipoSecundarioTest(EquipoSecundarioBaseTest):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.admin)

    def test_agrega_equipo_secundario(self):
        r = self.client.post(self.url, {"equipo_id": self.eq_b.pk})
        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertTrue(data["ok"], data)
        reg = JugadorEquipo.objects.get(jugador=self.jugador, equipo=self.eq_b)
        self.assertFalse(reg.es_principal)
        self.assertTrue(reg.activo)

    def test_no_modifica_datos_del_jugador(self):
        antes = (self.jugador.nombre, self.jugador.apellido,
                 self.jugador.curp, self.jugador.equipo_id,
                 self.jugador.fecha_nacimiento)
        self.client.post(self.url, {"equipo_id": self.eq_b.pk})
        self.jugador.refresh_from_db()
        self.assertEqual(
            (self.jugador.nombre, self.jugador.apellido,
             self.jugador.curp, self.jugador.equipo_id,
             self.jugador.fecha_nacimiento),
            antes,
        )

    def test_rechaza_equipo_principal(self):
        r = self.client.post(self.url, {"equipo_id": self.eq_a.pk})
        self.assertEqual(r.status_code, 400)
        self.assertFalse(r.json()["ok"])

    def test_rechaza_misma_categoria_que_principal(self):
        otro_a = Equipo.objects.create(nombre="EqA2", categoria=self.cat_a, activo=True)
        r = self.client.post(self.url, {"equipo_id": otro_a.pk})
        self.assertEqual(r.status_code, 400)
        self.assertFalse(r.json()["ok"])

    def test_rechaza_categoria_incompatible(self):
        r = self.client.post(self.url, {"equipo_id": self.eq_c.pk})
        self.assertEqual(r.status_code, 400)
        self.assertFalse(r.json()["ok"])

    def test_rechaza_por_edad(self):
        self.cat_b.edad_minima = 60
        self.cat_b.save()
        r = self.client.post(self.url, {"equipo_id": self.eq_b.pk})
        self.assertEqual(r.status_code, 400)
        self.assertIn("edad", r.json()["error"].lower())

    def test_es_idempotente(self):
        self.client.post(self.url, {"equipo_id": self.eq_b.pk})
        r = self.client.post(self.url, {"equipo_id": self.eq_b.pk})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(
            JugadorEquipo.objects.filter(jugador=self.jugador, equipo=self.eq_b).count(), 1
        )

    def test_rechaza_suspension_activa(self):
        SuspensionJugador.objects.create(
            jugador=self.jugador, categoria=self.cat_a, jornadas=2,
            fecha_inicio=date.today(),
        )
        r = self.client.post(self.url, {"equipo_id": self.eq_b.pk})
        self.assertEqual(r.status_code, 400)
        self.assertFalse(JugadorEquipo.objects.filter(
            jugador=self.jugador, equipo=self.eq_b).exists())

    def test_exige_csrf(self):
        client = self.client_class(enforce_csrf_checks=True)
        client.force_login(self.admin)
        r = client.post(self.url, {"equipo_id": self.eq_b.pk})
        self.assertEqual(r.status_code, 403)

    def test_anonimo_es_bloqueado(self):
        # El middleware de la app redirige al login antes de llegar a la vista.
        self.client.logout()
        r = self.client.post(self.url, {"equipo_id": self.eq_b.pk})
        self.assertIn(r.status_code, (302, 403))
        self.assertFalse(JugadorEquipo.objects.filter(
            jugador=self.jugador, equipo=self.eq_b).exists())

    def test_usuario_sin_permiso_es_403(self):
        usuario = get_user_model().objects.create_user(username="sinperm", password="p")
        self.client.force_login(usuario)
        r = self.client.post(self.url, {"equipo_id": self.eq_b.pk})
        self.assertEqual(r.status_code, 403)
        self.assertFalse(JugadorEquipo.objects.filter(
            jugador=self.jugador, equipo=self.eq_b).exists())

    def test_solo_post(self):
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 405)


class AllowedSecondaryCategoriesTest(TestCase):
    """_allowed_secondary_category_ids y _get_allowed_secondary_categories
    deben seguir dando el mismo resultado que antes del refactor."""

    def setUp(self):
        self.a = Categoria.objects.create(nombre="A", activo=True, curp_obligatoria=False)
        self.b = Categoria.objects.create(nombre="B", activo=True, curp_obligatoria=False)
        self.d = Categoria.objects.create(nombre="D", activo=True, curp_obligatoria=False)
        self.a.categorias_compatibles.add(self.b)
        self.b.categorias_compatibles.add(self.a)
        self.eq_a = Equipo.objects.create(nombre="EqA", categoria=self.a, activo=True)
        self.eq_b = Equipo.objects.create(nombre="EqB", categoria=self.b, activo=True)
        self.eq_d = Equipo.objects.create(nombre="EqD", categoria=self.d, activo=True)
        self.j = Jugador.objects.create(
            nombre="L", apellido="P", posicion="DEL", equipo=self.eq_a
        )

    def test_ambas_funciones_coinciden(self):
        from .forms import _allowed_secondary_category_ids, _get_allowed_secondary_categories

        casos = [[], [self.eq_b]]
        for eqs in casos:
            for eq in eqs:
                JugadorEquipo.objects.create(
                    jugador=self.j, equipo=eq, es_principal=False, activo=True
                )
            por_ids = set(_allowed_secondary_category_ids(self.eq_a, self.j))
            por_qs = set(_get_allowed_secondary_categories(self.eq_a, self.j)
                         .values_list("id", flat=True))
            self.assertEqual(por_ids, por_qs)
            JugadorEquipo.objects.filter(equipo__in=eqs).delete()

    def test_sin_secundario_permite_b(self):
        from .forms import _get_allowed_secondary_categories
        ids = set(_get_allowed_secondary_categories(self.eq_a, self.j)
                  .values_list("id", flat=True))
        self.assertEqual(ids, {self.b.pk})


class CambioEquipoSecundarioTest(EquipoSecundarioBaseTest):
    """Cambiar de equipo dentro de una categoría donde ya es secundario.

    Solo se permite si no hay temporada en curso, o si la hay con período de
    altas abierto y el jugador todavía no tiene participaciones en la categoría.
    """

    def setUp(self):
        super().setUp()
        self.eq_b2 = Equipo.objects.create(nombre="EqB2", categoria=self.cat_b, activo=True)
        self.registrar(self.jugador, self.eq_b)  # ya es secundario en B

    def abrir_temporada_b(self, con_altas=True, con_partido=False):
        t = Temporada.objects.create(
            categoria=self.cat_b, nombre="T-B",
            fecha_inicio=date.today() - timedelta(days=30),
            iniciada=True, activa=True, finalizada=False,
        )
        if con_altas:
            from .models import PeriodoAltas
            PeriodoAltas.objects.create(
                temporada=t, tipo="fechas", activo=True,
                fecha_inicio=date.today(),
                fecha_fin=date.today() + timedelta(days=30)
            )
        if con_partido:
            self._registrar_partido(t)
        return t

    def _registrar_partido(self, temporada):
        from .models import Campo, Jornada, Partido, JugadorPartido
        campo = Campo.objects.create(nombre="C", activo=True)
        jornada = Jornada.objects.create(
            numero=1, nombre="J1", temporada=temporada
        )
        partido = Partido.objects.create(
            temporada=temporada, jornada=jornada,
            equipo_local=self.eq_b, equipo_visitante=self.eq_b2, campo=campo,
            fecha_hora=datetime(2026, 1, 1, 12, 0), estado="JUG",
        )
        JugadorPartido.objects.create(
            jugador=self.jugador, partido=partido, equipo=self.eq_b
        )

    def test_cambia_de_equipo_sin_temporada_en_curso(self):
        self.client.force_login(self.admin)
        r = self.client.post(self.url, {"equipo_id": self.eq_b2.pk})
        self.assertEqual(r.status_code, 200, r.content)
        data = r.json()
        self.assertTrue(data["ok"], data)
        self.assertTrue(data["cambio"])
        # Solo queda el nuevo equipo en esa categoría.
        self.assertFalse(JugadorEquipo.objects.filter(
            jugador=self.jugador, equipo=self.eq_b).exists())
        reg = JugadorEquipo.objects.get(jugador=self.jugador, equipo=self.eq_b2)
        self.assertFalse(reg.es_principal)
        self.assertTrue(reg.activo)

    def test_cambia_con_altas_abiertas_y_sin_partidos(self):
        self.abrir_temporada_b(con_altas=True, con_partido=False)
        self.client.force_login(self.admin)
        r = self.client.post(self.url, {"equipo_id": self.eq_b2.pk})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertTrue(r.json()["cambio"])

    def test_no_cambia_sin_periodo_de_altas(self):
        self.abrir_temporada_b(con_altas=False)
        self.client.force_login(self.admin)
        r = self.client.post(self.url, {"equipo_id": self.eq_b2.pk})
        self.assertEqual(r.status_code, 400)
        self.assertIn("período de altas", r.json()["error"])
        self.assertTrue(JugadorEquipo.objects.filter(
            jugador=self.jugador, equipo=self.eq_b).exists())

    def test_no_cambia_si_ya_jugo_en_la_categoria(self):
        self.abrir_temporada_b(con_altas=True, con_partido=True)
        self.client.force_login(self.admin)
        r = self.client.post(self.url, {"equipo_id": self.eq_b2.pk})
        self.assertEqual(r.status_code, 400)
        self.assertIn("participaciones", r.json()["error"])
        self.assertTrue(JugadorEquipo.objects.filter(
            jugador=self.jugador, equipo=self.eq_b).exists())

    def test_repetir_el_mismo_equipo_es_no_op(self):
        # Doble clic: responde ok sin tocar nada, no es un error.
        self.client.force_login(self.admin)
        r = self.client.post(self.url, {"equipo_id": self.eq_b.pk})
        self.assertEqual(r.status_code, 200, r.content)
        data = r.json()
        self.assertTrue(data["ok"])
        self.assertTrue(data["sin_cambios"])
        self.assertEqual(JugadorEquipo.objects.filter(
            jugador=self.jugador, equipo=self.eq_b).count(), 1)

    def test_no_deja_dos_registros_en_la_misma_categoria(self):
        self.client.force_login(self.admin)
        self.client.post(self.url, {"equipo_id": self.eq_b2.pk})
        registros = JugadorEquipo.objects.filter(
            jugador=self.jugador, es_principal=False, equipo__categoria=self.cat_b
        )
        self.assertEqual(registros.count(), 1)

    def test_boton_visible_solo_por_cambio(self):
        # Ya tiene secundario en B, A~B compatible, y no hay categoria nueva:
        # el boton debe seguir visible porque puede cambiar de equipo en B.
        self.assertIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_boton_sigue_visible_para_poder_quitar_aunque_no_se_pueda_cambiar(self):
        # Temporada en curso, sin altas y con partido: no se puede ni alta ni
        # cambio, pero el secundario se debe poder quitar, asi que el boton
        # sigue visible.
        self.abrir_temporada_b(con_altas=False, con_partido=True)
        self.assertIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_boton_se_oculta_si_no_tiene_secundario_ni_se_puede_alta(self):
        # Sin secundarios y con la categoría destino cerrada por falta de altas.
        self.registros = JugadorEquipo.objects.filter(
            jugador=self.jugador, es_principal=False).delete()
        # Se cierra B para altas nuevas (temporada en curso sin periodo de altas)
        # y no hay mas categorias compatibles: no hay nada que hacer.
        t = self.abrir_temporada_b(con_altas=False)
        JugadorEquipo.objects.filter(jugador=self.jugador).delete()
        Jugador.objects.filter(pk=self.jugador.pk).update(equipo=self.eq_a)
        # Ahora solo queda la principal en A; A~B compatible pero B sin altas.
        self.assertNotIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_no_modifica_datos_del_jugador_al_cambiar(self):
        antes = (self.jugador.nombre, self.jugador.apellido,
                 self.jugador.curp, self.jugador.equipo_id,
                 self.jugador.fecha_nacimiento)
        self.client.force_login(self.admin)
        self.client.post(self.url, {"equipo_id": self.eq_b2.pk})
        self.jugador.refresh_from_db()
        self.assertEqual(
            (self.jugador.nombre, self.jugador.apellido,
             self.jugador.curp, self.jugador.equipo_id,
             self.jugador.fecha_nacimiento),
            antes,
        )


class QuitarEquipoSecundarioTest(EquipoSecundarioBaseTest):
    """Quitar (dar de baja) el equipo secundario de una categoría."""

    def setUp(self):
        super().setUp()
        self.client.force_login(self.admin)
        self.eq_b2 = Equipo.objects.create(nombre="EqB2", categoria=self.cat_b, activo=True)
        self.registro = self.registrar(self.jugador, self.eq_b)
        self.url = reverse("jugador_quitar_equipo_secundario", args=[self.jugador.pk])

    def test_quita_el_equipo_secundario(self):
        r = self.client.post(self.url, {"categoria_id": self.cat_b.pk})
        self.assertEqual(r.status_code, 200, r.content)
        data = r.json()
        self.assertTrue(data["ok"], data)
        self.assertTrue(data["cambio"])
        self.assertFalse(JugadorEquipo.objects.filter(pk=self.registro.pk).exists())
        self.assertEqual(data["categorias_secundarias"], [])

    def test_no_toca_el_equipo_principal_ni_los_datos(self):
        antes = (self.jugador.nombre, self.jugador.apellido,
                 self.jugador.curp, self.jugador.equipo_id)
        self.client.post(self.url, {"categoria_id": self.cat_b.pk})
        self.jugador.refresh_from_db()
        self.assertEqual(
            (self.jugador.nombre, self.jugador.apellido,
             self.jugador.curp, self.jugador.equipo_id),
            antes,
        )
        self.assertEqual(self.jugador.equipo_id, self.eq_a.pk)

    def test_conserva_el_registro_de_otra_categoria(self):
        self.registrar(self.jugador, self.eq_c)
        self.client.post(self.url, {"categoria_id": self.cat_b.pk})
        self.assertTrue(JugadorEquipo.objects.filter(
            jugador=self.jugador, equipo=self.eq_c, es_principal=False).exists())

    def test_es_idempotente(self):
        self.client.post(self.url, {"categoria_id": self.cat_b.pk})
        r = self.client.post(self.url, {"categoria_id": self.cat_b.pk})
        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertTrue(data["ok"])
        self.assertTrue(data["sin_cambios"])

    def test_no_puede_quitar_la_categoria_principal(self):
        r = self.client.post(self.url, {"categoria_id": self.cat_a.pk})
        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertTrue(data["sin_cambios"])
        # El equipo principal sigue intacto.
        self.assertEqual(self.jugador.equipo_id, self.eq_a.pk)
        self.assertFalse(JugadorEquipo.objects.filter(
            jugador=self.jugador, equipo=self.eq_a, es_principal=False).exists())

    def test_avisa_si_conserva_partidos(self):
        self._crear_partido()
        r = self.client.post(self.url, {"categoria_id": self.cat_b.pk})
        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertTrue(data["ok"])
        self.assertTrue(data["avisos"])
        self.assertIn("participacion", data["avisos"][0].lower())

    def test_exige_csrf(self):
        client = self.client_class(enforce_csrf_checks=True)
        client.force_login(self.admin)
        r = client.post(self.url, {"categoria_id": self.cat_b.pk})
        self.assertEqual(r.status_code, 403)

    def test_usuario_sin_permiso_es_403(self):
        usuario = get_user_model().objects.create_user(username="sinperm2", password="p")
        self.client.force_login(usuario)
        r = self.client.post(self.url, {"categoria_id": self.cat_b.pk})
        self.assertEqual(r.status_code, 403)
        self.assertTrue(JugadorEquipo.objects.filter(pk=self.registro.pk).exists())

    def test_anonimo_es_bloqueado(self):
        self.client.logout()
        r = self.client.post(self.url, {"categoria_id": self.cat_b.pk})
        self.assertIn(r.status_code, (302, 403))
        self.assertTrue(JugadorEquipo.objects.filter(pk=self.registro.pk).exists())

    def test_solo_post(self):
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 405)

    def test_categoria_invalida(self):
        r = self.client.post(self.url, {"categoria_id": "abc"})
        self.assertEqual(r.status_code, 400)

    def test_boton_visible_para_poder_quitar(self):
        # Ya tiene secundario: el botón debe verse aunque no pueda agregar más.
        self.assertIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_boton_se_oculta_si_no_tiene_secundarios_ni_destinos(self):
        JugadorEquipo.objects.filter(pk=self.registro.pk).delete()
        # Quedan A~B compatibles, así que sigue habiendo alta nueva posible.
        self.assertIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def _crear_partido(self):
        from .models import Campo, Jornada, Partido, JugadorPartido
        temp = Temporada.objects.create(
            categoria=self.cat_b, nombre="T-B",
            fecha_inicio=date.today() - timedelta(days=30),
            iniciada=True, activa=True, finalizada=False,
        )
        campo = Campo.objects.create(nombre="C", activo=True)
        jornada = Jornada.objects.create(numero=1, nombre="J1", temporada=temp)
        partido = Partido.objects.create(
            temporada=temp, jornada=jornada,
            equipo_local=self.eq_b, equipo_visitante=self.eq_b2, campo=campo,
            fecha_hora=datetime(2026, 1, 1, 12, 0), estado="JUG",
        )
        JugadorPartido.objects.create(
            jugador=self.jugador, partido=partido, equipo=self.eq_b
        )


class PeriodoAltasVencidoTest(EquipoSecundarioBaseTest):
    """Un PeriodoAltas con activo=True pero cuyo rango ya venció NO cuenta.

    Es el caso real que reporto el usuario: una categoría con temporada ya
    iniciada cuyo "periodo de altas" quedo marcado activo pero vencio. Debe
    quedar deshabilitada (no se puede dar de alta ni cambiar de equipo).
    """

    def setUp(self):
        super().setUp()
        self.t = Temporada.objects.create(
            categoria=self.cat_b, nombre="T-B",
            fecha_inicio=date.today() - timedelta(days=120),
            iniciada=True, activa=True, finalizada=False,
        )
        from .models import PeriodoAltas
        # activo=True pero el rango de fechas ya paso.
        self.periodo = PeriodoAltas.objects.create(
            temporada=self.t, tipo="fechas", activo=True,
            fecha_inicio=date.today() - timedelta(days=60),
            fecha_fin=date.today() - timedelta(days=30),
        )

    def test_el_periodo_vencido_no_cuenta_como_abierto(self):
        self.assertFalse(self.t.periodo_altas_activo())
        # El metodo ingenuo (solo el flag) si diria que si:
        self.assertTrue(self.t.periodos_altas.filter(activo=True).exists())

    def test_la_api_lo_informa_cerrado(self):
        self.client.force_login(self.admin)
        r = self.client.get(reverse("api_equipos_categoria"), {
            "equipo_id": self.eq_a.pk, "jugador_id": self.jugador.pk,
        })
        self.assertEqual(r.status_code, 200, r.content)
        data = r.json()
        self.assertFalse(data["periodos_abiertos"][str(self.cat_b.pk)])

    def test_el_endpoint_rechaza_el_alta(self):
        self.client.force_login(self.admin)
        r = self.client.post(
            reverse("jugador_agregar_equipo_secundario", args=[self.jugador.pk]),
            {"equipo_id": self.eq_b.pk},
        )
        self.assertEqual(r.status_code, 400)
        self.assertIn("per\u00edodo de altas", r.json()["error"])
        self.assertFalse(JugadorEquipo.objects.filter(
            jugador=self.jugador, equipo=self.eq_b).exists())

    def test_el_boton_no_se_muestra(self):
        # Unica categoria compatible y esta cerrada: no hay nada que hacer.
        self.assertNotIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))


class RegistroSecundarioInactivoTest(EquipoSecundarioBaseTest):
    """Un equipo secundario dado de baja (activo=False) debe volver a estar
    disponible como alta nueva.

    Caso real reportado en produccion: el jugador seguia apareciendo en el
    modal como "Ya registrado" en una categoria de la que ya lo habian dado
    de baja, y no podia volver a registrarse.
    """

    def setUp(self):
        super().setUp()
        self.reg = self.registrar(self.jugador, self.eq_b, es_principal=False, activo=False)
        self.client.force_login(self.admin)

    def test_el_modal_lo_ofrece_como_alta_nueva(self):
        r = self.client.get(reverse("api_equipos_categoria"), {
            "equipo_id": self.eq_a.pk, "jugador_id": self.jugador.pk,
        })
        self.assertEqual(r.status_code, 200, r.content)
        data = r.json()
        # B sigue siendo categoria compatible, pero NO esta registrada.
        self.assertIn(str(self.cat_b.pk), [str(c["id"]) for c in data["compatibles"]])
        self.assertNotIn(str(self.cat_b.pk), data["registros"])

    def test_se_puede_volver_a_registrar(self):
        r = self.client.post(
            reverse("jugador_agregar_equipo_secundario", args=[self.jugador.pk]),
            {"equipo_id": self.eq_b.pk},
        )
        self.assertEqual(r.status_code, 200, r.content)
        self.assertTrue(r.json()["ok"], r.json())
        self.reg.refresh_from_db()
        self.assertTrue(self.reg.activo)
        self.assertFalse(self.reg.es_principal)

    def test_el_boton_sigue_visible(self):
        # Sin secundarios activos pero con categoria compatible disponible.
        self.assertIn(self.jugador.pk, _jugadores_con_secundario_posible([self.jugador]))

    def test_el_registro_no_lo_lista_la_pantalla_de_registro(self):
        # Este es el sintoma: la pantalla de Registro solo muestra activos.
        eq = self.eq_b
        ids = JugadorEquipo.objects.filter(equipo=eq, activo=True).values_list(
            "jugador_id", flat=True)
        self.assertNotIn(self.jugador.pk, list(ids))


class SincronizacionEquipoPrincipalTest(TestCase):
    """Jugador.equipo y la inscripcion principal (JugadorEquipo) deben quedar
    siempre coherentes: quitar el equipo, por cualquier via, tiene que borrar
    la inscripcion; si queda, el jugador sigue figurando en el roster del
    equipo anterior y el reporte de Registro lo sigue listando.
    """

    def setUp(self):
        self.cat = Categoria.objects.create(nombre="SyncCat", es_principal=True)
        self.eq = Equipo.objects.create(nombre="SyncEq", categoria=self.cat, activo=True)
        self.eq2 = Equipo.objects.create(nombre="SyncEq2", categoria=self.cat, activo=True)
        self.jug = Jugador.objects.create(nombre="Sync", fecha_nacimiento=date(1995, 1, 1))

    def test_asignar_equipo_crea_la_inscripcion_principal(self):
        self.jug.equipo = self.eq
        self.jug.save()
        self.assertTrue(
            JugadorEquipo.objects.filter(
                jugador=self.jug, equipo=self.eq, es_principal=True
            ).exists()
        )

    def test_quitar_equipo_borra_la_inscripcion_principal(self):
        self.jug.equipo = self.eq
        self.jug.save()
        self.jug.equipo = None
        self.jug.save()
        self.assertFalse(JugadorEquipo.objects.filter(jugador=self.jug, es_principal=True).exists())

    def test_quitar_equipo_con_update_fields_tambien_borra(self):
        self.jug.equipo = self.eq
        self.jug.save()
        self.jug.equipo = None
        self.jug.save(update_fields=["equipo"])
        self.assertFalse(JugadorEquipo.objects.filter(jugador=self.jug, es_principal=True).exists())

    def test_no_toca_los_registros_secundarios(self):
        self.jug.equipo = self.eq
        self.jug.save()
        sec = JugadorEquipo.objects.create(jugador=self.jug, equipo=self.eq2, es_principal=False)
        self.jug.equipo = None
        self.jug.save()
        self.assertTrue(JugadorEquipo.objects.filter(pk=sec.pk, es_principal=False).exists())

    def test_update_fields_de_otro_campo_no_borra_la_inscripcion(self):
        self.jug.equipo = self.eq
        self.jug.save()
        # Salvar otro campo no debe desincronizar nada.
        self.jug.suspendido_pago = True
        self.jug.save(update_fields=["suspendido_pago"])
        self.assertTrue(
            JugadorEquipo.objects.filter(
                jugador=self.jug, equipo=self.eq, es_principal=True
            ).exists()
        )

    def test_reasignar_equipo_recrea_la_inscripcion(self):
        self.jug.equipo = self.eq
        self.jug.save()
        self.jug.equipo = None
        self.jug.save()
        self.jug.equipo = self.eq2
        self.jug.save()
        self.assertTrue(
            JugadorEquipo.objects.filter(
                jugador=self.jug, equipo=self.eq2, es_principal=True
            ).exists()
        )
        self.assertFalse(JugadorEquipo.objects.filter(jugador=self.jug, equipo=self.eq).exists())


class CedulaAcentosYLogosTest(TestCase):
    """La cedula no debe mostrar '?' donde iban acentos, y los logos de los
    equipos deben dibujarse grandes en la franja de cada equipo del PDF."""

    def setUp(self):
        import shutil
        import tempfile
        from django.contrib.auth import get_user_model
        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import override_settings
        from django.urls import reverse
        from PIL import Image as PILImage

        self.media = tempfile.mkdtemp(prefix="cedula_media_")
        ov = override_settings(MEDIA_ROOT=self.media)
        ov.enable()
        self.addCleanup(ov.disable)
        self.addCleanup(shutil.rmtree, self.media, ignore_errors=True)

        self.cat = Categoria.objects.create(nombre="CedulaLogoCat", dias_juego=["SAB"])
        self.campo = Campo.objects.create(nombre="CampoCedula", activo=True)
        self.local = Equipo.objects.create(nombre="SAN JUAN", categoria=self.cat, activo=True)
        self.visit = Equipo.objects.create(nombre="PABELLON", categoria=self.cat, activo=True)
        for eq in (self.local, self.visit):
            buf = BytesIO()
            PILImage.new("RGB", (200, 200), (200, 30, 30)).save(buf, format="PNG")
            eq.logo.save(
                f"{eq.nombre}.png",
                SimpleUploadedFile(f"{eq.nombre}.png", buf.getvalue(), content_type="image/png"),
                save=False,
            )
            eq.save()
        self.temp = Temporada.objects.create(
            categoria=self.cat, nombre="CedLogo", fecha_inicio=date(2026, 1, 3),
            tipo_rol="TODOS", vueltas=1, min_jugadores=0,
        )
        self.partido = Partido.objects.create(
            temporada=self.temp, equipo_local=self.local, equipo_visitante=self.visit,
            campo=self.campo, estado="PEND",
            fecha_hora=timezone.make_aware(datetime.combine(date(2026, 1, 10), time(15, 0))),
        )
        for eq, prefijo in ((self.local, "L"), (self.visit, "V")):
            for k in range(12):
                Jugador.objects.create(equipo=eq, nombre=f"{prefijo}{k} JUAN",
                                       apellido="PEREZ", dorsal=k + 1, activo=True)
        u = get_user_model().objects.create_superuser(username="cedlogo", password="p")
        self.client.force_login(u)
        self.pdf_url = reverse("reporte_cedula_arbitral_pdf", args=[self.partido.id])
        self.xlsx_url = reverse("reporte_cedula_arbitral_xlsx", args=[self.partido.id])

    def _pdf(self):
        import fitz
        resp = self.client.get(self.pdf_url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/pdf")
        doc = fitz.open(stream=resp.content, filetype="pdf")
        return doc, "".join(pg.get_text() for pg in doc)

    def test_pdf_no_tiene_interrogantes_en_la_cedula(self):
        _, texto = self._pdf()
        self.assertIn("CEDULA ARBITRAL", texto)
        self.assertIn("Arbitro:", texto)
        self.assertIn("Firma del Arbitro:", texto)
        for trozo in ("C?DULA", "C?dula", "?rbitro", "Firma del ?rbitro"):
            self.assertNotIn(trozo, texto)

    def test_pdf_dibuja_los_logos_de_equipo_mas_grandes(self):
        doc, texto = self._pdf()
        self.assertIn("LOCAL - SAN JUAN", texto)
        self.assertIn("VISITANTE - PABELLON", texto)
        altos = []
        for page in doc:
            for info in page.get_image_info():
                b = info["bbox"]
                altos.append(round(b[3] - b[1], 1))
        self.assertTrue(altos, "no se dibujo ningun logo en el PDF")
        # Antes era 22x22 dentro de una franja de 26; ahora 36x36 en una de 40.
        self.assertGreaterEqual(max(altos), 30.0, f"logos chicos: {altos}")

    def test_pdf_dibuja_mas_grande_el_logo_de_la_liga(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from PIL import Image as PILImage
        from league.models import ConfiguracionLiga
        cfg = ConfiguracionLiga.obtener()
        buf = BytesIO()
        PILImage.new("RGB", (200, 200), (20, 60, 120)).save(buf, format="PNG")
        cfg.logo.save(
            "liga.png",
            SimpleUploadedFile("liga.png", buf.getvalue(), content_type="image/png"),
            save=True,
        )
        doc, texto = self._pdf()
        self.assertIn("LOCAL - SAN JUAN", texto)
        altos = []
        for page in doc:
            for info in page.get_image_info():
                b = info["bbox"]
                altos.append(round(b[3] - b[1], 1))
        self.assertGreaterEqual(len(altos), 3, f"se esperaban 3 logos: {altos}")
        # El de la liga (46) tiene que ser el mas grande; los equipos van en 36.
        self.assertGreaterEqual(max(altos), 44.0, f"logo de liga chico: {altos}")

    def test_xlsx_titulo_y_datos_sin_interrogantes(self):
        import openpyxl
        resp = self.client.get(self.xlsx_url)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"].lower())
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws.title, "Cedula")
        joined = " ".join(
            str(ws.cell(row=r, column=c).value)
            for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)
            if ws.cell(row=r, column=c).value is not None
        )
        self.assertIn("Arbitro:", joined)
        self.assertNotIn("?", joined)


class InterrogantesPorAcentoTest(TestCase):
    """Varios archivos se escribieron con un encoder que cambio los acentos por
    '?' (Categor?a, suspensi?n, S?, pr?xima). Estos tests evitan que vuelva."""

    PATRON = re.compile(
        r"(?<![¿¡])[0-9A-Za-zÁÉÍÓÚÜÑáéíóúüñ]\?"
        r"|(?<![¿¡])\?[0-9A-Za-zÁÉÍÓÚÜÑáéíóúüñ]"
    )
    LETRAS = "0-9A-Za-zÁÉÍÓÚÜÑáéíóúüñ"
    # En HTML el '?' de cierre de pregunta es legitimo ("¿Va? </p>"), asi que
    # solo se marca si va dentro de una palabra, si abre un texto o si hace de
    # coma ("S?, enviar" era "Si, enviar").
    PATRON_HTML = re.compile(
        rf"[{LETRAS}]\?[{LETRAS}]"
        rf"|[{LETRAS}]\?[,;:]"
        rf"|[A-Za-zÁÉÍÓÚÑáéíóñ]\?\s"
        rf"|(?:>|\"|'|\(|\s)\?[{LETRAS}]"
    )

    @staticmethod
    def _raiz():
        return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

    def test_reports_views_no_tiene_interrogantes_en_los_textos(self):
        ruta = os.path.join(self._raiz(), "reports", "views.py")
        with open(ruta, encoding="utf-8-sig") as fh:
            arbol = ast.parse(fh.read())
        ofensores = []
        for nodo in ast.walk(arbol):
            if not isinstance(nodo, ast.Constant) or not isinstance(nodo.value, str):
                continue
            texto = nodo.value
            if "://" in texto or texto.strip() == "?":
                continue  # URLs y el '?' de respaldo para iniciales
            if self.PATRON.search(texto):
                ofensores.append("reports/views.py:%s -> %r" % (nodo.lineno, texto))
        self.assertEqual(ofensores, [], "\n".join(ofensores))

    def test_jornada_list_no_tiene_interrogantes_en_acentos(self):
        ruta = os.path.join(self._raiz(), "templates", "league", "jornada_list.html")
        with open(ruta, encoding="utf-8-sig") as fh:
            html = fh.read()
        # Se ignoran tags de Django y URLs: ahí el '?' es legitimo.
        html = re.sub(r"\{%.*?%\}", " ", html, flags=re.S)
        html = re.sub(r"\{\{.*?\}\}", " ", html, flags=re.S)
        html = re.sub(r"(?:href|src|action|data-href)\s*=\s*[\"'][^\"']*[\"']", " ", html, flags=re.I)
        ofensores = [l.strip() for l in html.splitlines() if self.PATRON_HTML.search(l)]
        self.assertEqual(ofensores, [], "\n".join(ofensores))


class ReporteSuscriptoresTest(TestCase):
    """El PDF de suscriptores se rompio cuando draw_header()gano un parametro
    (height): hay que llamar draw_header(p, w, h, extra) y los titulos van sin
    '?'."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        from django.urls import reverse

        from accounts.models import Rol
        from .models import SuscripcionEmail

        self.rol = Rol.objects.create(
            nombre="ReporteSubs", permisos={"reporte_suscriptores": True}
        )
        self.user = get_user_model().objects.create_user(
            username="subs", password="p", rol=self.rol
        )
        self.cat = Categoria.objects.create(nombre="Libre", dias_juego=["SAB"])
        self.sub = SuscripcionEmail.objects.create(
            email="socio@ejemplo.com", usuario=self.user, activo=True
        )
        self.sub.categorias.add(self.cat)
        self.client.force_login(self.user)
        self.pdf_url = reverse("reporte_suscriptores_pdf")
        self.xlsx_url = reverse("reporte_suscriptores_xlsx")

    def test_pdf_se_genera_sin_interrogantes(self):
        import fitz
        resp = self.client.get(self.pdf_url, {"categoria": self.cat.id})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/pdf")
        texto = "".join(
            pg.get_text() for pg in fitz.open(stream=resp.content, filetype="pdf")
        )
        for trozo in ("Email", "Usuario", "Categorias", "Roles", "Estadisticas",
                      "Activo", "Creado", "SI", "socio@ejemplo.com"):
            self.assertIn(trozo, texto)
        self.assertNotIn("?", texto)

    def test_xlsx_titulos_sin_interrogantes(self):
        import openpyxl
        resp = self.client.get(self.xlsx_url, {"categoria": self.cat.id})
        self.assertEqual(resp.status_code, 200)
        ws = openpyxl.load_workbook(BytesIO(resp.content)).active
        joined = " ".join(
            str(ws.cell(row=r, column=c).value)
            for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)
            if ws.cell(row=r, column=c).value is not None
        )
        self.assertIn("Categorias", joined)
        self.assertIn("Estadisticas", joined)
        self.assertIn("SI", joined)
        self.assertNotIn("?", joined)
